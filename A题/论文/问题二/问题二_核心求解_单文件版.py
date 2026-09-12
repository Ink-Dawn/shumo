# -*- coding: utf-8 -*-
"""
2026 A题 · 第二问核心求解（修正版，单文件）
========================================
发给队友：本文件 + 题目原始“附件1.xlsx”，不需要原来的q2文件夹。

建议使用Python 3.12。安装依赖：
    python -m pip install numpy==1.26.0 scipy==1.16.1 openpyxl==3.1.5

把附件放在本文件旁边，然后运行：
    python 问题二_核心求解_单文件版.py
也可以指定附件和输出位置：
    python 问题二_核心求解_单文件版.py --input "附件1.xlsx" --output "计算结果"

常用修改：直接修改下方Config；或通过命令行覆盖：
    --N 2560 --h 25 --hm 8e-7 --end 10800
快速检查：--self-test；短时试跑：--end 600
保存每个内部节点的完整场：--full-field（文件较大）。

默认参数对应已核验的2560格正式修正版。完整3小时通常需要数分钟。
默认保存每秒、21个题目半径的未舍入解；默认网格联立计算全部2561个节点。
本次单文件整理未重新运行完整三小时计算。
输出：result2.xlsx、temperature_required.csv（表三）、moisture_required.csv
（表四）、solution.npz（未舍入结果）、balance_check.csv、run_info.json。
短时试跑的表格仅包含实际完成的时间和正文时刻。

方法：圆柱径向有限体积 + 表面加密sinh网格 + 非线性势差水分通量 + BDF。
热方程采用B(C)*T_t，不是[B(C)*T]_t。温度状态用°C，D公式使用K。
h、hm沿用第一问是模型假设；环境水分列作为有效Robin边界驱动。
此文件负责正式核心计算与导出；独立有限元和整套敏感性研究另有原项目记录。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import platform
from dataclasses import asdict, dataclass
from functools import lru_cache
from pathlib import Path
from time import perf_counter

try:
    import numpy as np
    import scipy
    import openpyxl
    from numpy.polynomial.legendre import leggauss
    from scipy.integrate import solve_ivp
    from scipy.sparse import diags, kron
    from openpyxl import Workbook, load_workbook
except ModuleNotFoundError as error:
    raise SystemExit(
        f"缺少依赖 {error.name}。请先运行：python -m pip install numpy scipy openpyxl"
    ) from error


# 1. 修改模型参数通常只需要改这一段；默认值是正式修正版。
@dataclass(frozen=True)
class Config:
    N: int = 2560                  # 径向区间数，节点数=N+1
    R: float = 0.02               # 圆柱半径，m
    L: float = 0.25               # 圆柱长度，m
    h: float = 25.0               # 换热系数，W/(m²·K)
    hm: float = 8e-7              # 水分交换系数，m/s
    T0: float = 28.0              # 初始温度，°C
    C0: float = 2.55              # 初始含水率，kg/kg
    beta: float = 2.5             # sinh网格表面加密程度；0表示均匀网格
    gauss: int = 5                # 非线性通量积分点数
    method: str = 'BDF'           # 可改为Radau进行时间方法对照
    rtol: float = 1e-13
    atol_T: float = 1e-14
    atol_C: float = 1e-15
    max_step: float = 5.0         # 最大内部时间步，s；不等于输出间隔
    end: int = 10800              # 总时长，s；每1秒输出


def check_config(p: Config) -> None:
    if p.N < 2 or p.end < 1 or p.gauss < 1:
        raise ValueError('N至少为2，end与gauss应为正整数。')
    numbers = [p.R, p.L, p.h, p.hm, p.T0, p.C0, p.beta,
               p.rtol, p.atol_T, p.atol_C, p.max_step]
    if not np.isfinite(numbers).all():
        raise ValueError('模型参数必须为有限数值。')
    if min(p.R, p.L, p.C0, p.rtol, p.atol_T, p.atol_C, p.max_step) <= 0:
        raise ValueError('尺寸、初始含水率、时间步和误差容限应大于0。')
    if p.T0 <= -273.15 or p.h < 0 or p.hm < 0 or p.beta < 0:
        raise ValueError('温度须高于绝对零度，h、hm与beta不能为负。')
    if p.method not in ('BDF', 'Radau'):
        raise ValueError('method请选择BDF或Radau。')


# 2. 独立读取原始附件，不读取或硬编码之前算出的答案。
@dataclass
class Environment:
    time: np.ndarray
    temperature: np.ndarray
    moisture: np.ndarray
    source: str = ''
    sha256: str = ''

    def at(self, t):
        if np.any(np.asarray(t) < self.time[0]) or np.any(np.asarray(t) > self.time[-1]):
            raise ValueError('环境时刻超出附件范围，禁止外推。')
        return np.interp(t, self.time, self.temperature), np.interp(t, self.time, self.moisture)


def find_input(path: Path | None) -> Path:
    if path is not None:
        candidate = path.expanduser().resolve()
        if not candidate.is_file():
            raise FileNotFoundError(f'附件不存在：{candidate}')
        return candidate
    base = Path(__file__).resolve().parent
    for candidate in (base/'附件1.xlsx', base/'附件'/'附件1.xlsx',
                      base.parent/'附件'/'附件1.xlsx', Path.cwd()/'附件1.xlsx'):
        if candidate.is_file():
            return candidate.resolve()
    raise FileNotFoundError('找不到附件1.xlsx。请放在本脚本旁，或用 --input 指定完整路径。')


def read_environment(path: Path, end: int = 10800) -> Environment:
    path = Path(path).resolve()
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        # 附件第一行为表头；前三列依次为秒、环境温度、环境水分。
        rows = list(wb.active.values)[1:]
        a = np.asarray([row[:3] for row in rows if any(v is not None for v in row)], float)
    finally:
        wb.close()
    if a.ndim != 2 or a.shape[1] != 3 or len(a) < 2 or not np.isfinite(a).all():
        raise ValueError('附件应有三列完整数值：时间(s)、环境温度(°C)、环境水分。')
    if a[0, 0] != 0 or np.any(np.diff(a[:, 0]) <= 0) or a[-1, 0] < end:
        raise ValueError('环境数据须从0开始、时间严格递增，且覆盖计算终点。')
    if np.any(a[:, 2] <= 0):
        raise ValueError('环境水分必须为正。')
    return Environment(*a.T, str(path), hashlib.sha256(path.read_bytes()).hexdigest())


# 3. 附录物性函数。D的第二个参数必须是绝对温度K。
def rho(C):
    return 650.0 + 128.0*C


def cp(C):
    return 1450.0 + 2736.0*C/(C+1.0)


def k(C):
    return 0.21 + 0.38*C/(C+1.0)


def B(C):
    return rho(C)*cp(C)


def A(T_K):
    return 2.4e-3*np.exp(-3850.0/T_K)


def D(C, T_K):
    return A(T_K)*np.exp(-0.45/C)


# 4. 圆环网格：保留圆心和表面，使用真实圆柱体积、界面面积。
@dataclass
class Grid:
    r: np.ndarray
    rf: np.ndarray
    volumes: np.ndarray
    areas: np.ndarray


def make_grid(p: Config) -> Grid:
    x = np.linspace(0.0, 1.0, p.N+1)
    r = p.R*(1-np.sinh(p.beta*(1-x))/np.sinh(p.beta)) if p.beta else p.R*x
    rf = np.r_[0.0, (r[:-1]+r[1:])/2, p.R]
    return Grid(r, rf, np.pi*p.L*np.diff(rf**2), 2*np.pi*p.L*rf)


def interpolate(values, r, target):
    """将时间×节点矩阵线性插值到所需半径；不平滑、不提前舍入。"""
    hi = np.clip(np.searchsorted(r, target, side='right'), 1, len(r)-1)
    w = (target-r[hi-1])/(r[hi]-r[hi-1])
    return values[:, hi-1]*(1-w) + values[:, hi]*w


# 5. 界面通量；正号统一表示朝圆柱外侧流出。
@lru_cache(None)
def rule(order):
    s, w = leggauss(order)
    return (s+1)/2, w/2


def delta_phi(left, right, order=5):
    """计算∫[C_left,C_right] exp(-0.45/C)dC，保留通量非线性。"""
    d = right-left
    s, w = rule(order)
    return d*np.sum(w[:, None]*np.exp(-0.45/(left[None, :]+s[:, None]*d[None, :])), axis=0)


def fluxes(T, C, g, p, env_T, env_C, frozen=None):
    ft, fc = np.zeros(len(T)+1), np.zeros(len(C)+1)
    if frozen is None:
        kval = k(C)
        kface = 2*kval[:-1]*kval[1:]/(kval[:-1]+kval[1:])
        fc[1:-1] = -g.areas[1:-1]*A((T[:-1]+T[1:])/2+273.15)*delta_phi(C[:-1], C[1:], p.gauss)/np.diff(g.r)
    else:
        # 仅供另行进行冻结物性辅助检查；默认正式计算不走此分支。
        _, kface, diffusion = frozen
        fc[1:-1] = -g.areas[1:-1]*diffusion*np.diff(C)/np.diff(g.r)
    ft[1:-1] = -g.areas[1:-1]*kface*np.diff(T)/np.diff(g.r)
    ft[-1] = g.areas[-1]*p.h*(T[-1]-env_T)
    fc[-1] = g.areas[-1]*p.hm*(C[-1]-env_C)
    # 中心面积为0，ft[0]和fc[0]自然为0。
    return ft, fc


# 6. 联立方程右端；状态交错存储为[T0,C0,T1,C1,...]。
def sparsity(N):
    blocks = diags([np.ones(N), np.ones(N+1), np.ones(N)], [-1, 0, 1], format='csr')
    return kron(blocks, np.ones((2, 2)), format='csc')


def make_rhs(p, g, env, frozen=None):
    def rhs(t, y):
        T, C = y[0::2], y[1::2]
        if not np.isfinite(y).all() or np.any(C <= 0) or np.any(T <= -273.15):
            raise ValueError(f'非物理试探状态 t={t}: minC={C.min()}, minT={T.min()}')
        ft, fc = fluxes(T, C, g, p, *env.at(t), frozen=frozen)
        dy = np.empty_like(y)
        capacity = B(C) if frozen is None else frozen[0]
        dy[0::2] = -np.diff(ft)/(capacity*g.volumes)
        dy[1::2] = -np.diff(fc)/g.volumes
        return dy
    return rhs


# 7. 每个环境折点重启BDF，整秒只用于记录；物理模型和原修正版一致。
@dataclass
class Result:
    time: np.ndarray
    r: np.ndarray
    T: np.ndarray
    C: np.ndarray
    average_C: np.ndarray
    inventory: np.ndarray
    surface_integral: np.ndarray
    integral_check: np.ndarray
    heat_balance: np.ndarray
    stats: dict
    grid: Grid


def solve(p, env, full=False, frozen=None, progress=True):
    check_config(p)
    start = perf_counter()
    g = make_grid(p)
    rhs = make_rhs(p, g, env, frozen)
    y = np.empty(2*(p.N+1)); y[0::2] = p.T0; y[1::2] = p.C0
    atol = np.empty_like(y); atol[0::2] = p.atol_T; atol[1::2] = p.atol_C
    time = np.arange(p.end+1, dtype=float)
    r = g.r if full else np.linspace(0, p.R, 21)
    Tout, Cout = np.empty((len(time), len(r))), np.empty((len(time), len(r)))
    inventory = np.empty(len(time)); hbal = np.empty(len(time))
    integrals = {n: np.zeros(len(time)) for n in (4, 8)}
    cumulative = {4: 0.0, 8: 0.0}
    stats = dict(nfev=0, njev=0, nlu=0, segments=0, config=asdict(p), success=True,
                 min_temperature_C=p.T0, min_moisture=p.C0)
    knots = np.r_[env.time[(env.time >= 0) & (env.time < p.end)], float(p.end)]
    for j, (left, right) in enumerate(zip(knots[:-1], knots[1:])):
        sol = solve_ivp(rhs, (left, right), y, method=p.method, rtol=p.rtol, atol=atol,
                        max_step=p.max_step, jac_sparsity=sparsity(p.N), dense_output=True)
        if not sol.success or abs(sol.t[-1]-right) > 1e-9:
            raise RuntimeError(f'第{j}段未到达终点: {sol.message}')
        ids = np.where((time >= left) & (time <= right))[0]
        z = sol.sol(time[ids]); T, C = z[0::2].T, z[1::2].T
        if not np.isfinite(z).all() or (C.size and np.min(C) <= 0):
            raise RuntimeError('输出含非有限值或非正含水率。')
        Tout[ids] = T if full else interpolate(T, g.r, r)
        Cout[ids] = C if full else interpolate(C, g.r, r)
        inventory[ids] = C@g.volumes
        if C.size:
            stats['min_temperature_C'] = min(stats['min_temperature_C'], float(T.min()))
            stats['min_moisture'] = min(stats['min_moisture'], float(C.min()))
        # 检查给定热方程的瞬时储热率，不将它解释为额外总内能模型。
        for idx, tt, zz in zip(ids, time[ids], z.T):
            capacity = B(zz[1::2]) if frozen is None else frozen[0]
            outward = g.areas[-1]*p.h*(zz[-2]-env.at(tt)[0])
            hbal[idx] = np.sum(g.volumes*capacity*rhs(tt, zz)[0::2])+outward
        # 独立积分表面通量：比较4点和8点求积，检查有效含水率收支。
        breaks = np.unique(np.r_[sol.t, time[ids]])
        lefts, width = breaks[:-1], np.diff(breaks)
        for order in (4, 8):
            s, w = rule(order)
            tq = lefts[:, None]+width[:, None]*s
            surface_C = sol.sol(tq.ravel())[-1].reshape(tq.shape)
            flux = g.areas[-1]*p.hm*(surface_C-env.at(tq)[1])
            values = cumulative[order]+np.r_[0., np.cumsum(width*np.sum(w*flux, axis=1))]
            integrals[order][ids] = values[np.searchsorted(breaks, time[ids])]
            cumulative[order] = float(values[-1])
        y = sol.y[:, -1]
        for key in ('nfev', 'njev', 'nlu'):
            stats[key] += getattr(sol, key)
        stats['segments'] += 1
        if progress and (j % 30 == 0 or right == p.end):
            print(f'{p.method}  N={p.N}  已计算 {right:.0f}/{p.end} 秒', flush=True)
    stats['runtime_s'] = perf_counter()-start
    return Result(time, r, Tout, Cout, inventory/g.volumes.sum(), inventory,
                  integrals[8], integrals[8]-integrals[4], hbal, stats, g)


# 8. 导出：所有表格来自本次未舍入解，不读取预存答案。
def verify_excel(path, time, radii, T, C):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if wb.sheetnames != ['温度', '水分浓度']:
            raise AssertionError('工作表名称错误。')
        for sheet, expected in zip(wb, (T, C)):
            rows = list(sheet.values)
            if len(rows) != len(time) or len(rows[0]) != len(radii)+1:
                raise AssertionError('Excel行列数错误。')
            header_radii = np.asarray([float(f'{x:.14g}') for x in radii*100])
            if not np.array_equal(np.asarray(rows[0][1:], float), header_radii):
                raise AssertionError('Excel半径表头不一致。')
            values = np.asarray(rows[1:], float)
            if not np.isfinite(values).all() or not np.array_equal(values[:, 0], time[1:]):
                raise AssertionError('Excel时间列或数值不正确。')
            if not np.array_equal(values[:, 1:], np.round(expected[1:], 4)):
                raise AssertionError('Excel与未舍入结果的四位小数不一致。')
    finally:
        wb.close()
    return dict(status='PASS', data_rows_per_sheet=len(time)-1,
                radii_per_sheet=len(radii), numeric_result_cells=2*(len(time)-1)*len(radii))


def export_results(z, p, env, output, full=False):
    output = Path(output).resolve(); output.mkdir(parents=True, exist_ok=True)
    radii = np.linspace(0, p.R, 21)
    T = z.T if np.array_equal(z.r, radii) else interpolate(z.T, z.r, radii)
    C = z.C if np.array_equal(z.r, radii) else interpolate(z.C, z.r, radii)
    wb = Workbook(); wb.remove(wb.active)
    body_times = np.arange(1800, p.end+1, 1800, dtype=int)
    # 小时表仍取圆心、R/4、R/2、3R/4、表面；R=.02m时即题目五个半径。
    headers = ['time_h']+[f'r={x:g}cm' for x in radii[::5]*100]
    for sheet_name, values, filename in [('温度', T, 'temperature_required.csv'), ('水分浓度', C, 'moisture_required.csv')]:
        with (output/filename).open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.writer(stream); writer.writerow(headers)
            for t in body_times:
                writer.writerow([f'{t/3600:.4f}']+[f'{x:.4f}' for x in values[t, ::5]])
        ws = wb.create_sheet(sheet_name)
        # 表头去除二进制浮点尾差；计算仍使用原始半径，不提前舍入解。
        # 默认R对应原题0、0.1、…、2 cm，也支持修改R后的实际半径。
        ws.append(['时间(s)/半径(cm)']+[float(f'{x:.14g}') for x in radii*100])
        for t, row in zip(z.time[1:], np.round(values[1:], 4)):
            ws.append([int(t)]+row.tolist())
        ws.freeze_panes = 'B2'; ws.column_dimensions['A'].width = 20
        for row in ws.iter_rows(min_row=2, min_col=2):
            for cell in row: cell.number_format = '0.0000'
    excel_path = output/'result2.xlsx'; wb.save(excel_path); wb.close()
    excel_check = verify_excel(excel_path, z.time, radii, T, C)
    raw_path = output/'solution.npz'
    np.savez_compressed(raw_path, time_s=z.time, radius_m=z.r, temperature_C=z.T, moisture=z.C,
                        volume_average_moisture=z.average_C, environment_temperature_C=env.at(z.time)[0],
                        environment_moisture=env.at(z.time)[1], saved_all_internal_nodes=np.asarray(full))
    residual = z.inventory-z.inventory[0]+z.surface_integral
    with (output/'balance_check.csv').open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(['time_s','effective_inventory','surface_integral','relative_residual','Gauss8_minus_Gauss4','heat_rate_residual_W'])
        writer.writerows(zip(z.time,z.inventory,z.surface_integral,residual/z.inventory[0],z.integral_check,z.heat_balance))
    balance_check = dict(moisture_relative=float(np.max(abs(residual))/z.inventory[0]),
                         flux_quadrature_relative=float(np.max(abs(z.integral_check))/z.inventory[0]),
                         heat_rate_residual_W=float(np.max(abs(z.heat_balance))))
    info = dict(status='SOLVED AND EXPORTED', config=asdict(p),
                input_path=env.source,input_sha256=env.sha256,
                script_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
                python=platform.python_version(),numpy=np.__version__,scipy=scipy.__version__,openpyxl=openpyxl.__version__,
                scope='Full nonlinear FVM core calculation; full convergence/FEM/sensitivity study is separate',
                complete_three_hour_run=p.end==10800,full_internal_field_saved=full,
                statistics=z.stats,excel_check=excel_check,balance_check=balance_check,
                final=dict(time_s=float(z.time[-1]),temperature_center_C=float(z.T[-1,0]),temperature_surface_C=float(z.T[-1,-1]),
                           moisture_center=float(z.C[-1,0]),moisture_surface=float(z.C[-1,-1]),volume_average_moisture=float(z.average_C[-1])),
                files={name:hashlib.sha256((output/name).read_bytes()).hexdigest() for name in
                       ('result2.xlsx','temperature_required.csv','moisture_required.csv','solution.npz','balance_check.csv')})
    (output/'run_info.json').write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf-8')
    return info


def self_test():
    """快速检查单位、几何、零通量平衡与非线性势差积分。"""
    p = Config(N=8, end=60); check_config(p); g = make_grid(p)
    assert abs(g.volumes.sum()/(np.pi*p.R**2*p.L)-1) < 1e-12
    assert g.r[0]==0 and g.r[-1]==p.R and np.all(g.volumes>0)
    assert np.isclose(D(2.55,301.15),2.4e-3*np.exp(-.45/2.55)*np.exp(-3850/301.15),rtol=1e-14,atol=0)
    env = Environment(np.array([0.,60.]),np.array([28.,28.]),np.array([2.55,2.55]))
    y = np.tile([28.,2.55],p.N+1)
    assert np.max(abs(make_rhs(p,g,env)(0,y))) < 1e-12
    c = np.linspace(.2,2.55,10)
    assert np.array_equal(delta_phi(c,c),np.zeros(10))
    assert np.max(abs(delta_phi(c,c+.01,5)-delta_phi(c,c+.01,8))) < 1e-12
    print('自检通过：物性温标、圆柱体积、平衡态和通量积分。',flush=True)


def main():
    parser = argparse.ArgumentParser(description='第二问2560格修正版：单文件核心求解与结果导出',
                                     formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    parser.add_argument('--input',type=Path,help='附件1.xlsx路径；省略时自动寻找')
    parser.add_argument('--output',type=Path,default=Path(__file__).resolve().parent/'问题二_单文件结果',help='结果文件夹')
    for name in ('N','end','gauss'):
        parser.add_argument('--'+name,type=int,default=getattr(Config(),name))
    for name in ('R','L','h','hm','T0','C0','beta','rtol','atol_T','atol_C','max_step'):
        parser.add_argument('--'+name,type=float,default=getattr(Config(),name))
    parser.add_argument('--method',choices=['BDF','Radau'],default=Config().method)
    parser.add_argument('--full-field',action='store_true',help='保存全部内部节点；默认只保存21个输出半径')
    parser.add_argument('--self-test',action='store_true',help='仅做快速基础检查，不读取附件、不生成正式结果')
    args = parser.parse_args()
    if args.self_test:
        self_test(); return
    p = Config(**{name:getattr(args,name) for name in Config.__dataclass_fields__}); check_config(p)
    path = find_input(args.input); env = read_environment(path,p.end)
    output = args.output.expanduser().resolve()
    # 避免误将输入Excel选作输出目标；正常默认输出不接触题目附件。
    if path == output/'result2.xlsx':
        raise ValueError('输出result2.xlsx不能覆盖输入附件。请修改--output。')
    print(f'读取附件：{path}\n输出目录：{output}',flush=True)
    result = solve(p,env,full=args.full_field)
    info = export_results(result,p,env,output,full=args.full_field)
    print(f"完成：{info['excel_check']['numeric_result_cells']}个Excel结果已回读核对。",flush=True)
    print('末态中心/表面温度：'+f'{result.T[-1,0]:.4f} / {result.T[-1,-1]:.4f} °C',flush=True)
    print('末态中心/表面含水率：'+f'{result.C[-1,0]:.4f} / {result.C[-1,-1]:.4f} kg/kg',flush=True)
    print(f'结果文件：{output / "result2.xlsx"}',flush=True)


if __name__ == '__main__':
    main()
