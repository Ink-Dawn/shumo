# -*- coding: utf-8 -*-
"""
2026 A题 · 第三问核心求解（单文件交接版）
=======================================
发给队友：本文件 + 题目原始“附件1.xlsx”，不需要原来的 q3 文件夹。
建议 Python 3.12；依赖：
    python -m pip install numpy==1.26.0 scipy==1.16.1 openpyxl==3.1.5

把附件放在本文件旁边，执行：
    python 问题三_核心求解_单文件版.py
指定附件、输出位置和运行名称：
    python 问题三_核心求解_单文件版.py --input "附件1.xlsx" --output "计算结果" --run-id trial
修改参数：直接改下方 Config，或提供 --config 参数.json；命令行参数最后覆盖。
例如：--N 1280 --method Radau --scenario T_plus --run-id trial2
只检查环境、配置和输入（不积分）：--check-only
只导出已有完成运行：--export-only --output "计算结果" --run-id trial
可选 --template 指定题目 result3.xlsx 模板；未提供时创建同名 Sheet1 和完整表头。

默认输出到本文件旁的“单文件结果/runs/baseline”：
    result3.xlsx；tables/table5.csv、table5.md、table5_unrounded.npz；
    raw/chunk_*.npz（全部内部节点的温湿场）；diagnostics.csv；
    event.json；critical_state.npz；event_neighborhood.npz；checkpoints/；manifest.json。
未到阈值则标记 not_reached，保存原始场和检查点，不伪造临界行或正式结果表。

默认参数对应第三问正式 N=640、S60 配置；本次封装不重新进行完整烘干计算。
第二问正式 N=2560，与第三问默认配置的前3h四位小数并非处处相同。
保留原求解内核：圆柱径向有限体积、表面加密网格、完整非线性势差通量、
BDF/Radau、自适应时间步、4h单侧环境切换、全节点事件与真实后验积分。
时间以秒计，温度状态为°C（扩散系数公式换算K），半径以m计，含水率kg/kg。
S60表示4h后保持末60min时间均值；该环境延拓是建模假设。
Excel首列秒，21个半径列默认0—2cm；60s主体序列，末行为精确临界事件。
四位小数只用于显示；判定、检查点和原始场均保留未舍入值。

恢复示例（保留同一输出目录中的父运行）：
    python 问题三_核心求解_单文件版.py --output "计算结果" --run-id resumed --resume "计算结果/runs/trial/checkpoints/latest.npz"
未提供 --config 时，恢复自动读取检查点配置。需 --branch 才可调整后段时间参数；
长期情景只能在4h检查点切换。运行ID不能覆盖不同参数/代码版本的旧运行。
只支持本单文件版本生成的检查点；修改代码后需重新计算，避免混用版本。
整套空间收敛、制造解、情景扫描、独立有限元及论文制图不在此核心交接文件中。
修改方法或参数后应由使用者重新做相应验证；求解器成功不等于物理模型已验证。
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import platform
from collections import deque
from copy import copy
from dataclasses import asdict, dataclass, replace
from datetime import datetime, timezone
from functools import lru_cache
from pathlib import Path
from time import perf_counter

# 与原运行保持单线程数值库设置；在导入 NumPy / SciPy 之前生效。
for _thread_key in ('OMP_NUM_THREADS', 'OPENBLAS_NUM_THREADS', 'MKL_NUM_THREADS', 'NUMEXPR_NUM_THREADS'):
    os.environ[_thread_key] = '1'

try:
    import numpy as np
    import scipy
    from numpy.polynomial.legendre import leggauss
    from scipy.integrate import BDF, Radau, solve_ivp
    from scipy.optimize import brentq
    from scipy.sparse import diags, kron
    from openpyxl import Workbook, load_workbook
except ModuleNotFoundError as error:
    raise SystemExit(f'缺少依赖 {error.name}。请执行：python -m pip install numpy scipy openpyxl') from error

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_OUTPUT = BASE_DIR / '单文件结果'
SCENARIOS = ('S0', 'S30', 'S60', 'S90', 'T_minus', 'T_plus', 'Ce_minus', 'Ce_plus')
ORIGINAL_CORE_SHA256 = 'ff1cf55434c3cfff2194e1801820647c18527bb459de6af0e22be0a58a96d400'


# 1. 常用修改区：以下默认数值来自 configs/accepted_config.json。



@dataclass(frozen=True)
class Config:
    N: int = 640                  # 径向区间数，内部节点数=N+1
    R: float = 0.02               # 固定圆柱半径，m
    L: float = 0.25               # 固定长度，m
    h: float = 25.0               # 换热系数，W/(m²·K)
    hm: float = 8e-7              # 水分交换系数，m/s
    T0: float = 28.0              # 初始温度，°C
    C0: float = 2.55              # 初始含水率，kg/kg
    beta: float = 2.5             # 表面加密；0表示均匀网格
    gauss: int = 5                # 非线性势差通量的Gauss阶数
    method: str = 'BDF'           # BDF 或 Radau
    rtol: float = 1e-10           # 4h后相对容差
    atol_T: float = 1e-11         # 温度绝对容差
    atol_C: float = 1e-12         # 含水率绝对容差
    max_step: float = 150.0       # 后期最大内部步长，s
    early_rtol: float = 1e-10     # 前4h参数；与输出时间间隔不同
    early_atol_T: float = 1e-11
    early_atol_C: float = 1e-12
    early_max_step: float = 10.0
    end: float = 432000.0         # 搜索上限120h；不是预先指定的答案
    sample_s: float = 60.0        # 保存全节点场的间隔；必须整除60s
    checkpoint_s: float = 21600.0 # 后期默认每6h持久化并重启
    Ccrit: float = 0.15           # 全域含水率临界阈值
    scenario: str = 'S60'
    post_s: float = 10.0          # 事件后真实积分时长，至少1s

    def check(self):
        if not isinstance(self.N, int) or isinstance(self.N, bool) or self.N < 2:
            raise ValueError('N必须为至少2的整数。')
        if not isinstance(self.gauss, int) or self.gauss < 1:
            raise ValueError('gauss必须为正整数。')
        if self.method not in ('BDF', 'Radau') or self.scenario not in SCENARIOS:
            raise ValueError('method或scenario取值不支持。')
        if any(not np.isfinite(v) for v in asdict(self).values() if isinstance(v, (int, float))):
            raise ValueError('所有数值参数必须为有限值。')
        positive = (self.R, self.L, self.h, self.hm, self.C0, self.rtol, self.atol_T,
                    self.atol_C, self.max_step, self.early_rtol, self.early_atol_T,
                    self.early_atol_C, self.early_max_step, self.end, self.sample_s,
                    self.checkpoint_s, self.Ccrit)
        if min(positive) <= 0 or self.beta < 0 or self.T0 <= -273.15 or self.post_s < 1:
            raise ValueError('尺寸、容差、步长等须为正；T0须高于绝对零度，post_s至少1。')
        if self.C0 <= self.Ccrit:
            raise ValueError('本入口搜索下降穿越，初始C0必须大于Ccrit。')
        if self.sample_s > 60 or not np.isclose(60 / self.sample_s, round(60 / self.sample_s), rtol=0, atol=1e-9):
            raise ValueError('sample_s必须为60s的正约数，以保证规定导出时刻已保存。')
        return self

    @classmethod
    def read(cls, path):
        return cls(**json.loads(Path(path).read_text('utf-8-sig'))).check()


# 2. 输入定位；附件是唯一必需的数据文件，模板可选。
def find_input(path=None):
    if path is not None:
        result = Path(path).expanduser().resolve()
        if not result.is_file():
            raise FileNotFoundError(f'找不到附件：{result}')
        return result
    candidates = (BASE_DIR / '附件1.xlsx', BASE_DIR / 'inputs/附件1.xlsx',
                  BASE_DIR / '附件/附件1.xlsx', BASE_DIR.parent / '附件/附件1.xlsx',
                  Path.cwd() / '附件1.xlsx')
    for result in candidates:
        if result.is_file():
            return result.resolve()
    raise FileNotFoundError('请把附件1.xlsx放在本文件旁，或用 --input 指定路径。')


def checked_run_id(run_id):
    if not run_id or run_id in ('.', '..') or any(c in run_id for c in '<>:"/\\|?*') or run_id.endswith((' ', '.')):
        raise ValueError('run-id须为单个有效文件夹名称，不能包含路径分隔符。')
    return run_id


def output_path(output_root=None):
    return Path(output_root if output_root is not None else DEFAULT_OUTPUT).expanduser().resolve()



def window_mean(t,v,window):
    left=t[-1]-window
    if window<=0 or left<t[0]: raise ValueError('窗口超出观测范围')
    x=np.r_[left,t[(t>left)&(t<t[-1])],t[-1]]
    return float(np.trapz(np.interp(x,t,v),x)/window)


@dataclass
class Environment:
    time:np.ndarray
    temperature:np.ndarray
    moisture:np.ndarray
    plateau_T:float
    plateau_C:float
    scenario:str
    side:str='observed'

    def at(self,t):
        x=np.asarray(t)
        if self.side=='observed':
            if np.any(x<self.time[0]) or np.any(x>self.time[-1]): raise ValueError('观测边界禁止外推')
            return np.interp(t,self.time,self.temperature),np.interp(t,self.time,self.moisture)
        if self.side!='plateau' or np.any(x<self.time[-1]-1e-8): raise ValueError('平台边界时刻错误')
        return np.full_like(x,self.plateau_T,dtype=float),np.full_like(x,self.plateau_C,dtype=float)

    def values(self,t):
        x=np.asarray(t)
        return (np.where(x<=self.time[-1],np.interp(x,self.time,self.temperature),self.plateau_T),
                np.where(x<=self.time[-1],np.interp(x,self.time,self.moisture),self.plateau_C))


def read_environment(scenario='S60',path=None):
    path=find_input(path)
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        rows=list(wb.active.values)
        a=np.array([r[:3] for r in rows[1:] if any(v is not None for v in r)],float)
    finally: wb.close()
    if a.ndim!=2 or a.shape[1]!=3 or not np.isfinite(a).all(): raise ValueError('附件三列应完整且有限')
    t,T,C=a.T
    if t[0]!=0 or t[-1]!=14400 or np.any(np.diff(t)<=0) or np.any(C<=0) or np.any(T<=-273.15):
        raise ValueError('附件时间/单位/物理范围与方案不符')
    windows={'S30':1800,'S60':3600,'S90':5400}
    if scenario=='S0': pt,pc=float(T[-1]),float(C[-1])
    else:
        w=windows.get(scenario,3600)
        pt,pc=window_mean(t,T,w),window_mean(t,C,w)
        shifts={'T_minus':(-1,0),'T_plus':(1,0),'Ce_minus':(0,-.005),'Ce_plus':(0,.005)}
        if scenario not in windows and scenario not in shifts: raise ValueError(f'未知情景{scenario}')
        dt,dc=shifts.get(scenario,(0,0)); pt+=dt; pc+=dc
    if pt<=-273.15 or pc<=0: raise ValueError('平台不满足物理正性')
    return Environment(t,T,C,pt,pc,scenario)


def statistics():
    env=read_environment(); rows=[]
    for name,w in [('S30',1800),('S60',3600),('S90',5400)]:
        mask=env.time>=env.time[-1]-w
        x=(env.time[mask]-env.time[-1])/3600
        rows.append(dict(scenario=name,window_min=w/60,T_C=window_mean(env.time,env.temperature,w),
                    Ce=window_mean(env.time,env.moisture,w),
                    T_slope_C_per_h=float(np.polyfit(x,env.temperature[mask],1)[0]),
                    Ce_slope_per_h=float(np.polyfit(x,env.moisture[mask],1)[0])))
    return rows


# 3. 物性与圆柱网格：与原问题三物理内核相同。


def rho(C):
    return 650.0 + 128.0 * C


def cp(C):
    return 1450.0 + 2736.0 * C / (C + 1.0)


def k(C):
    return 0.21 + 0.38 * C / (C + 1.0)


def B(C):
    return rho(C) * cp(C)


def A(T_K):
    return 2.4e-3 * np.exp(-3850.0 / T_K)


def D(C, T_K):
    return A(T_K) * np.exp(-0.45 / C)


@dataclass
class Grid:
    r: np.ndarray
    rf: np.ndarray
    volumes: np.ndarray
    areas: np.ndarray


def make_grid(p):
    x = np.linspace(0.0, 1.0, p.N + 1)
    r = p.R * (1 - np.sinh(p.beta * (1 - x)) / np.sinh(p.beta)) if p.beta else p.R*x
    rf = np.r_[0.0, (r[:-1]+r[1:])/2, p.R]
    return Grid(r, rf, np.pi*p.L*np.diff(rf**2), 2*np.pi*p.L*rf)


def interpolate(values, r, target):
    """时间×节点矩阵，线性插值至固定输出半径。"""
    hi = np.clip(np.searchsorted(r, target, side='right'), 1, len(r)-1)
    w = (target-r[hi-1])/(r[hi]-r[hi-1])
    return values[:, hi-1]*(1-w) + values[:, hi]*w


# 4. 非线性通量与半离散方程：状态交错排列T0,C0,T1,C1,...。


@lru_cache(None)
def rule(order):
    s, w = leggauss(order)
    return (s+1)/2, w/2


def delta_phi(left, right, order=5):
    d = right-left
    s, w = rule(order)
    return d * np.sum(w[:, None]*np.exp(-0.45/(left[None, :]+s[:, None]*d[None, :])), axis=0)


def fluxes(T, C, g, p, env_T, env_C, frozen=None):
    ft, fc = np.zeros(len(T)+1), np.zeros(len(C)+1)
    if frozen is None:
        kval = k(C)
        kface = 2*kval[:-1]*kval[1:]/(kval[:-1]+kval[1:])
        fc[1:-1] = -g.areas[1:-1] * A((T[:-1]+T[1:])/2+273.15) * delta_phi(C[:-1], C[1:], p.gauss)/np.diff(g.r)
    else:
        _, kface, diffusion = frozen
        fc[1:-1] = -g.areas[1:-1]*diffusion*np.diff(C)/np.diff(g.r)
    ft[1:-1] = -g.areas[1:-1]*kface*np.diff(T)/np.diff(g.r)
    ft[-1] = g.areas[-1]*p.h*(T[-1]-env_T)
    fc[-1] = g.areas[-1]*p.hm*(C[-1]-env_C)
    return ft, fc


def sparsity(N):
    blocks = diags([np.ones(N), np.ones(N+1), np.ones(N)], [-1, 0, 1], format='csr')
    return kron(blocks, np.ones((2, 2)), format='csc')


def make_rhs(p, g, env, frozen=None):
    def rhs(t, y):
        T, C = y[0::2], y[1::2]
        if not np.isfinite(y).all() or np.any(C <= 0) or np.any(T <= -273.15):
            raise ValueError(f'非物理状态 t={t}: minC={C.min()}, minT={T.min()}')
        ft, fc = fluxes(T, C, g, p, *env.at(t), frozen=frozen)
        dy = np.empty_like(y)
        capacity = B(C) if frozen is None else frozen[0]
        dy[0::2] = -np.diff(ft)/(capacity*g.volumes)
        dy[1::2] = -np.diff(fc)/g.volumes
        return dy
    return rhs


# 5. 全域事件判定：全部内部C节点的最大值，Brent定位与显式括区间细化。


def event_value(y,threshold=.15): return float(np.max(y[1::2])-threshold)


def first_bracket(times,values):
    for a,b,ga,gb in zip(times[:-1],times[1:],values[:-1],values[1:]):
        if ga>=0 and gb<0: return float(a),float(b)
    return None


def locate(dense,bracket,threshold):
    fun=lambda t:event_value(dense(t),threshold)
    a,b=bracket
    root=brentq(fun,a,b,xtol=1e-7,rtol=4*np.finfo(float).eps)
    left,right=a,b
    while right-left>1e-5:
        mid=(left+right)/2
        if fun(mid)>0: left=mid
        else: right=mid
    return dict(t_critical_s=float(root),t_critical_h=float(root/3600),Ccrit=threshold,
                g_at_root=fun(root),initial_bracket_s=[a,b],refined_bracket_s=[left,right],
                refined_bracket_width_s=right-left,root_method='Brent + explicit bisection bracket',
                root_tolerance_s=1e-7,event_direction=-1)


def diagnostic(t,y,g,p,env):
    C=y[1::2]; maximum=float(C.max()); tied=np.flatnonzero(maximum-C<=1e-10)
    inventory=float(C@g.volumes)
    return dict(time_s=float(t),Cmax=maximum,Ccenter=float(C[0]),Csurface=float(C[-1]),
                Cmean=inventory/float(g.volumes.sum()),min_C=float(C.min()),
                rmax_m=float(g.r[np.argmax(C)]),tied_maxima=len(tied),
                tied_rmin_m=float(g.r[tied[0]]),tied_rmax_m=float(g.r[tied[-1]]),
                Emono=float(np.maximum(np.diff(C),0).max()),
                center_max_gap=maximum-float(C[0]),inventory=inventory,
                surface_flux=float(g.areas[-1]*p.hm*(C[-1]-env.at(t)[1])),g=maximum-p.Ccrit)


# 6. 版本、原始场与检查点；无需原项目快照。


def sha(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for b in iter(lambda:f.read(8*1024*1024),b''): h.update(b)
    return h.hexdigest()


def save_json(path,value):
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    temp=path.with_suffix(path.suffix+'.tmp')
    temp.write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
    temp.replace(path)


def now(): return datetime.now(timezone.utc).isoformat()


def core_hash():
    """检查点绑定整个单文件版本，不依赖原项目目录。"""
    return sha(__file__)


def checkpoint(path,t,y,p,run_id,integrals,input_hash,code_hash):
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    temp=path.with_name(path.stem+'.tmp.npz')
    np.savez_compressed(temp,time_s=t,state=y,integrals=integrals,
        metadata=json.dumps(dict(config=p,run_id=run_id,input_hash=input_hash,code_hash=code_hash)))
    temp.replace(path)


def read_checkpoint(path):
    with np.load(path) as z:
        return float(z['time_s']),z['state'].copy(),z['integrals'].copy(),json.loads(str(z['metadata']))


def load_raw(run_id, output_root=None):
    folder=output_path(output_root)/'runs'/checked_run_id(run_id)
    index=json.loads((folder/'raw/index.json').read_text('utf-8'))
    blocks=[]
    if index.get('parent_run'):
        parent=load_raw(index['parent_run'],output_root); mask=parent['time_s']<index['start_s']-1e-7
        blocks.append({k:v[mask] for k,v in parent.items() if k!='radius_internal_m'})
    radius=None
    for item in index['chunks']:
        with np.load(folder/item['path']) as z:
            radius=z['radius_internal_m'].copy()
            blocks.append({k:z[k].copy() for k in ['time_s','temperature_C','moisture']})
    if not blocks: raise ValueError('没有有效原始场')
    answer={k:np.concatenate([b[k] for b in blocks],axis=0) for k in blocks[-1]}
    answer['radius_internal_m']=radius
    assert np.all(np.diff(answer['time_s'])>0)
    return answer


def memory_mb():
    try:
        import psutil
        return psutil.Process(os.getpid()).memory_info().rss/1024**2
    except ImportError: return None


# 7. 长时推进；保持原接受步、通量求积、环境切换及真实后验逻辑。


def run(p:Config,run_id='baseline',role='standalone',resume=None,branch=False,*,input_path=None,output_root=None):
    p.check(); root=output_path(output_root); run_id=checked_run_id(run_id)
    source=find_input(input_path); folder=root/'runs'/run_id
    mpath=folder/'manifest.json'
    code=core_hash(); ihash=sha(source)
    if mpath.exists():
        old=json.loads(mpath.read_text('utf-8'))
        if old['config']==asdict(p) and old['source_code_hash']==code and old['input_hash']==ihash and old['status'] in ('completed','not_reached'):
            print(f'复用已完成运行 {run_id}',flush=True); return old
        raise RuntimeError(f'运行ID已存在且不能复用：{run_id}；保留原记录，另选运行ID或指定恢复入口')
    (folder/'raw').mkdir(parents=True,exist_ok=True)
    (folder/'checkpoints').mkdir(exist_ok=True)
    env=read_environment(p.scenario,source); tb=float(env.time[-1]); grid=make_grid(p)
    parent=None; t=0.; y=np.tile([p.T0,p.C0],p.N+1); totals=np.zeros(2)
    if resume:
        t,y,totals,meta=read_checkpoint(resume); parent=meta['run_id']
        if meta['input_hash']!=ihash or meta['code_hash']!=code: raise ValueError('检查点输入或核心版本不一致')
        allowed={'rtol','atol_T','atol_C','max_step','method','end','sample_s','checkpoint_s','post_s'} if branch else {'end'}
        if branch and abs(t-tb)<1e-8: allowed.add('scenario')
        differences={k for k in asdict(p) if asdict(p)[k]!=meta['config'][k]}
        if differences-allowed: raise ValueError(f'检查点配置不可复用：{differences-allowed}')
    if not 0 <= t < p.end:
        raise ValueError('恢复时刻必须早于本次搜索上限end。')
    if y.shape != (2*(p.N+1),) or not np.isfinite(y).all() or np.any(y[1::2] <= 0):
        raise ValueError('初始/恢复状态的维数或含水率无效。')
    if event_value(y,p.Ccrit) <= 0:
        raise ValueError('恢复状态已达阈值；请从更早的检查点恢复或直接导出原事件。')
    if parent:
        checked_run_id(parent)
        parent_folder=root/'runs'/parent
        if not (parent_folder/'raw/index.json').is_file():
            raise ValueError('恢复需要同一输出目录中的父运行raw记录，请保留整个输出目录。')
        parent_manifest=json.loads((parent_folder/'manifest.json').read_text('utf-8'))
        if parent_manifest['source_code_hash'] != code or parent_manifest['input_hash'] != ihash:
            raise ValueError('父运行与检查点的输入/代码版本不一致。')
    start_t=t
    index=dict(start_s=t,parent_run=parent,chunks=[],units=dict(time_s='s',radius_internal_m='m',temperature_C='degC',moisture='kg/kg'),sampling_s=p.sample_s)
    save_json(folder/'raw/index.json',index)
    manifest=dict(run_id=run_id,role=role,status='running',config=asdict(p),source_code_hash=code,input_hash=ihash,
        source_file=Path(__file__).name,input_path=str(source),original_project_core_hash=ORIGINAL_CORE_SHA256,
        config_hash=__import__('hashlib').sha256(json.dumps(asdict(p),sort_keys=True).encode()).hexdigest(),
        started_utc=now(),pid=os.getpid(),python=platform.python_version(),numpy=np.__version__,scipy=scipy.__version__,
        threads={k:os.environ.get(k) for k in ['OMP_NUM_THREADS','OPENBLAS_NUM_THREADS','MKL_NUM_THREADS']},
        start_s=t,parent_run=parent,resume_path=str(resume) if resume else None,
        environment=dict(T_C=env.plateau_T,Ce=env.plateau_C,reachable_asymptotically=env.plateau_C<p.Ccrit,
            boundary_switch='observed left at 14400 s; plateau right restart, continuous state'),
        restart_policy='observed knots, then absolute checkpoint multiples',command_arguments=['python',str(Path(__file__).resolve()),'--input',str(source),'--output',str(root),'--run-id',run_id,'--config',str(folder/'config.json')])
    save_json(mpath,manifest); save_json(folder/'config.json',asdict(p))
    # 初态也是真正可恢复的检查点：首段失败不能指向尚未生成的文件。
    np.savez_compressed(folder/'raw/chunk_0000.npz',time_s=np.array([t]),radius_internal_m=grid.r,
        temperature_C=y[0::2][None,:],moisture=y[1::2][None,:],volumes=grid.volumes,areas=grid.areas,
        environment_temperature_C=env.values(np.array([t]))[0],environment_moisture=env.values(np.array([t]))[1])
    index['chunks'].append(dict(path='raw/chunk_0000.npz',first_s=t,last_s=t,shape=[1,p.N+1],sha256=sha(folder/'raw/chunk_0000.npz')))
    save_json(folder/'raw/index.json',index)
    checkpoint(folder/'checkpoints/latest.npz',t,y,asdict(p),run_id,totals,ihash,code)
    wall=perf_counter(); steps=[]; stats=dict(nfev=0,njev=0,nlu=0,accepted_steps=0,segments=0)
    initial_inventory=p.C0*float(grid.volumes.sum())
    recent=deque(); last_saved=t; peak=memory_mb() or 0.; event=None
    maxima=dict(moisture_balance_relative=0.,flux_quadrature_relative=0.,heat_rate_W=0.,Emono=0.,center_max_gap=0.,min_C=float(y[1::2].min()))
    writer=None; logfile=(folder/'diagnostics.csv').open('w',encoding='utf-8-sig',newline='')
    segment_records=[]; last_print=wall
    try:
        while t<p.end-1e-9 and event is None:
            if t<tb-1e-8:
                side='observed'; stop=min(float(env.time[np.searchsorted(env.time,t+1e-8)]),p.end)
                rtol,at,ac,maxstep=p.early_rtol,p.early_atol_T,p.early_atol_C,p.early_max_step
            else:
                side='plateau'; stop=min((np.floor((t+1e-7)/p.checkpoint_s)+1)*p.checkpoint_s,p.end)
                rtol,at,ac,maxstep=p.rtol,p.atol_T,p.atol_C,p.max_step
            boundary=replace(env,side=side); rhs=make_rhs(p,grid,boundary)
            atol=np.tile([at,ac],p.N+1)
            solver=({'BDF':BDF,'Radau':Radau}[p.method])(rhs,t,y,stop,rtol=rtol,atol=atol,
                max_step=maxstep,jac_sparsity=sparsity(p.N))
            segment_start=t; segment_start_g=event_value(y,p.Ccrit); rows=[]; states=[]
            if last_saved<start_t:
                rows.append(t); states.append(y.copy()); last_saved=t
            if writer is None:
                row=diagnostic(t,y,grid,p,boundary); row.update(flux_integral_4=float(totals[0]),flux_integral_8=float(totals[1]),
                    balance_relative=float((row['inventory']-initial_inventory+totals[1])/initial_inventory),step_s=0.,interior_min_g=row['g'])
                writer=csv.DictWriter(logfile,fieldnames=list(row)); writer.writeheader(); writer.writerow(row)
            while solver.status=='running':
                old_t=solver.t; old_y=solver.y.copy(); message=solver.step()
                if solver.status=='failed': raise RuntimeError(f't={solver.t}: {message}')
                right=float(solver.t); dense=solver.dense_output(); dt=right-old_t
                if not np.isfinite(solver.y).all() or np.min(solver.y[1::2])<=0: raise RuntimeError('非物理接受状态')
                probe=np.linspace(old_t,right,5); trial=dense(probe)
                gv=np.max(trial[1::2],axis=0)-p.Ccrit
                bracket=first_bracket(probe,gv)
                effective_right=right
                if bracket is not None:
                    event=locate(dense,bracket,p.Ccrit); effective_right=event['t_critical_s']
                    event['accepted_step_bracket_s']=[old_t,right]
                    event['accepted_step_endpoint_g']=[float(gv[0]),float(gv[-1])]
                if gv[-1]<.02:
                    recent.append((old_t,right,dense))
                    while len(recent)>1 and recent[0][1]<effective_right-120: recent.popleft()
                for j,order in enumerate((4,8)):
                    s,w=rule(order); tq=old_t+(effective_right-old_t)*s
                    cs=dense(tq)[-1]
                    totals[j]+=(effective_right-old_t)*float(np.sum(w*grid.areas[-1]*p.hm*(cs-boundary.at(tq)[1])))
                times=np.arange((np.floor((last_saved+1e-7)/p.sample_s)+1)*p.sample_s,effective_right+1e-8,p.sample_s)
                times=times[times>last_saved+1e-7]
                if len(times):
                    rows.extend(times.tolist()); states.extend(dense(times).T); last_saved=float(times[-1])
                if event is not None and effective_right>last_saved+1e-7:
                    rows.append(effective_right); states.append(dense(effective_right)); last_saved=effective_right
                yy=dense(effective_right) if event else solver.y
                row=diagnostic(effective_right,yy,grid,p,boundary)
                balance=(row['inventory']-initial_inventory+totals[1])/initial_inventory
                row.update(flux_integral_4=float(totals[0]),flux_integral_8=float(totals[1]),balance_relative=float(balance),step_s=dt,interior_min_g=float(gv.min()))
                writer.writerow(row)
                maxima['moisture_balance_relative']=max(maxima['moisture_balance_relative'],abs(balance))
                maxima['flux_quadrature_relative']=max(maxima['flux_quadrature_relative'],abs(totals[1]-totals[0])/initial_inventory)
                heat=np.sum(grid.volumes*B(yy[1::2])*rhs(effective_right,yy)[0::2])+grid.areas[-1]*p.h*(yy[-2]-boundary.at(effective_right)[0])
                maxima['heat_rate_W']=max(maxima['heat_rate_W'],abs(float(heat)))
                for k in ('Emono','center_max_gap'): maxima[k]=max(maxima[k],row[k])
                maxima['min_C']=min(maxima['min_C'],row['min_C'])
                steps.append(dt); stats['accepted_steps']+=1
                if event: break
            t=event['t_critical_s'] if event else float(solver.t)
            y=dense(t).copy() if event else solver.y.copy()
            for key in ('nfev','njev','nlu'): stats[key]+=getattr(solver,key)
            stats['segments']+=1
            segment_records.append(dict(start_s=segment_start,end_s=t,boundary=side,g_start=segment_start_g,g_end=event_value(y,p.Ccrit),status=solver.status))
            if rows:
                values=np.array(states); path=f'raw/chunk_{len(index["chunks"]):04d}.npz'
                np.savez_compressed(folder/path,time_s=np.array(rows),radius_internal_m=grid.r,temperature_C=values[:,0::2],moisture=values[:,1::2],
                    volumes=grid.volumes,areas=grid.areas,environment_temperature_C=boundary.at(np.array(rows))[0],environment_moisture=boundary.at(np.array(rows))[1])
                index['chunks'].append(dict(path=path,first_s=rows[0],last_s=rows[-1],shape=list(values[:,0::2].shape),sha256=sha(folder/path)))
                save_json(folder/'raw/index.json',index)
            checkpoint(folder/'checkpoints/latest.npz',t,y,asdict(p),run_id,totals,ihash,code)
            if abs(t-tb)<1e-7 or abs(t/p.checkpoint_s-round(t/p.checkpoint_s))<1e-8 or event:
                checkpoint(folder/f'checkpoints/state_{t:.6f}.npz',t,y,asdict(p),run_id,totals,ihash,code)
            peak=max(peak,memory_mb() or 0.)
            save_json(folder/'progress.json',dict(time_s=t,Cmax=float(y[1::2].max()),elapsed_s=perf_counter()-wall,latest_checkpoint='checkpoints/latest.npz',statistics=stats))
            logfile.flush()
            if perf_counter()-last_print>15 or side=='plateau' or event:
                print(f'{run_id}: {t/3600:.3f} h, Cmax={y[1::2].max():.7f}, 已用{perf_counter()-wall:.1f}s',flush=True); last_print=perf_counter()
        if event:
            np.savez_compressed(folder/'critical_state.npz',time_s=t,state=y,radius_internal_m=grid.r)
            post=solve_ivp(rhs,(t,t+p.post_s),y,method=p.method,rtol=min(p.rtol,1e-10),
                atol=np.tile([min(p.atol_T,1e-11),min(p.atol_C,1e-12)],p.N+1),
                max_step=1.,jac_sparsity=sparsity(p.N),dense_output=True)
            if not post.success: raise RuntimeError('事件后真实积分失败')
            deltas=np.array([1.,2.,5.,10.]); deltas=deltas[deltas<=p.post_s]
            pre=[]
            for delta in deltas:
                tt=t-delta
                match=next((s for a,b,s in recent if a-1e-9<=tt<=b+1e-9),None)
                if match is None: raise RuntimeError('缺少事件前有效连续解')
                pre.append(match(tt))
            pre=np.array(pre); post_states=post.sol(t+deltas).T
            np.savez_compressed(folder/'event_neighborhood.npz',deltas_s=deltas,pre_time_s=t-deltas,pre_states=pre,post_time_s=t+deltas,post_states=post_states)
            gpre=np.max(pre[:,1::2],axis=1)-p.Ccrit; gpost=np.max(post_states[:,1::2],axis=1)-p.Ccrit
            event.update(pre_time_s=float(t-deltas[-1]),post_time_s=float(t+deltas[-1]),delta_s=float(deltas[-1]),g_pre=float(gpre[-1]),g_post=float(gpost[-1]),
                post_state_source='true_integration',strict_post_pass=bool(np.all(gpre>0) and np.all(gpost<0)),
                slope_estimates=[dict(delta_s=float(d),slope_per_s=float((b-a)/(2*d))) for d,a,b in zip(deltas,gpre,gpost)],
                controlling_nodes=np.flatnonzero(y[1::2].max()-y[1::2]<=1e-10).tolist(),center_max_gap=float(y[1::2].max()-y[1]),
                first_crossing_evidence='all accepted endpoints plus quarter-step samples; earliest sampled downward bracket; numerical evidence, not a continuous theorem',
                critical_state_path='critical_state.npz',post_state_path='event_neighborhood.npz',numerical_validation_status='needs_validation')
            if not event['strict_post_pass']: raise RuntimeError('临界前后严格穿越检查失败')
            save_json(folder/'event.json',event)
        manifest.update(status='completed' if event else 'not_reached',finished_utc=now(),end_s=t,event_found=event is not None,
            t_critical_s=event['t_critical_s'] if event else None,runtime_s=perf_counter()-wall,peak_observed_memory_mb=peak,
            statistics=stats,step_statistics_s=dict(min=float(min(steps)),median=float(np.median(steps)),max=float(max(steps)),p95=float(np.percentile(steps,95))),
            diagnostics=maxima,segment_records=segment_records,raw_index_sha256=sha(folder/'raw/index.json'))
        save_json(mpath,manifest); return manifest
    except Exception as error:
        manifest.update(status='failed',error=repr(error),failed_utc=now(),runtime_s=perf_counter()-wall,latest_valid_checkpoint='checkpoints/latest.npz')
        save_json(mpath,manifest); raise
    finally: logfile.close()


def regular_and_event(end,interval,tolerance=1e-6):
    times=np.arange(interval,end+tolerance,interval)
    if len(times) and abs(times[-1]-end)<=tolerance: times[-1]=end
    else: times=np.r_[times[times<end],end]
    return times


def extract_times(raw,times):
    ids=np.searchsorted(raw['time_s'],times)
    ids=np.minimum(ids,len(raw['time_s'])-1)
    assert np.all(abs(raw['time_s'][ids]-times)<1e-6), '所需时刻必须来自实际保存场'
    return raw['moisture'][ids]



# 8. 导出：从同一运行的未舍入场提取，只给真正找到的临界事件追加末行。
def export_run(run_id='baseline', output_root=None, template=None):
    root = output_path(output_root)
    folder = root / 'runs' / checked_run_id(run_id)
    manifest = json.loads((folder / 'manifest.json').read_text('utf-8'))
    if manifest['status'] != 'completed' or not manifest['event_found']:
        raise ValueError('该运行未完成临界事件，不能导出正式表5/result3.xlsx。')
    if manifest['source_code_hash'] != core_hash():
        raise ValueError('当前单文件版本与该运行不同，不能混用导出。')
    event = json.loads((folder / 'event.json').read_text('utf-8'))
    if not event['strict_post_pass'] or event['post_state_source'] != 'true_integration':
        raise ValueError('事件缺少真实积分的严格后验。')
    end = event['t_critical_s']
    raw = load_raw(run_id, root)
    p = Config(**manifest['config']).check()
    radii = np.arange(21) * (p.R / 20)
    times = regular_and_event(end, 60.)
    values = interpolate(extract_times(raw, times), raw['radius_internal_m'], radii)
    if template is None:
        wb = Workbook()
        wb.active.title = 'Sheet1'
    else:
        template = Path(template).expanduser().resolve()
        if template == (folder / 'result3.xlsx').resolve():
            raise ValueError('模板不能与输出文件是同一路径。')
        wb = load_workbook(template)
        if wb.sheetnames != ['Sheet1']:
            wb.close()
            raise ValueError('模板须仅含题目约定的Sheet1。')
    ws = wb['Sheet1']
    style = copy(ws.cell(2, 2)._style)
    for row in ws:
        for cell in row:
            cell.value = None
    ws.cell(1, 1, '时间(s) / 到中心距离(cm)')
    for j, r in enumerate(radii * 100, 2):
        ws.cell(1, j, float(r))
    for i, (t, cs) in enumerate(zip(times, values), 2):
        ws.cell(i, 1, float(t)).number_format = '0.0000'
        for j, c in enumerate(cs, 2):
            cell = ws.cell(i, j, float(c))
            cell._style = copy(style)
            cell.number_format = '0.0000'
    ws.freeze_panes = 'B2'
    ws.column_dimensions['A'].width = 26
    # 模板中的预留格式行/列也会影响max_row/max_column；短程结果要裁去空尾部。
    if ws.max_row > len(times) + 1:
        ws.delete_rows(len(times) + 2, ws.max_row - len(times) - 1)
    if ws.max_column > 22:
        ws.delete_cols(23, ws.max_column - 22)
    path = folder / 'result3.xlsx'
    wb.save(path)
    wb.close()

    # 回读核对储存值。这里不把“导出通过”称作物理/算法验证通过。
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb['Sheet1'].values)
        stored = np.array(rows[1:], float)
        if stored.shape != (len(times), 22) or not np.isfinite(stored).all():
            raise ValueError('Excel尺寸或有限值检查失败。')
        if not np.allclose(np.array(rows[0][1:], float), radii*100, rtol=0, atol=1e-12):
            raise ValueError('Excel半径表头检查失败。')
        if np.max(abs(stored[:, 0] - times)) > 1e-8 or np.any(np.diff(stored[:, 0]) <= 0):
            raise ValueError('Excel时间序列检查失败。')
        if np.max(abs(stored[:, 1:] - values)) > 1e-14:
            raise ValueError('Excel储存值回读不一致。')
        for row in wb['Sheet1'].iter_rows(min_row=2):
            if any(cell.number_format != '0.0000' for cell in row):
                raise ValueError('Excel四位小数显示格式不一致。')
    finally:
        wb.close()

    body_t = regular_and_event(end, 21600.)
    body_r = np.arange(5) * (p.R / 4)
    body_c = interpolate(extract_times(raw, body_t), raw['radius_internal_m'], body_r)
    tables = folder / 'tables'
    tables.mkdir(exist_ok=True)
    with (tables / 'table5.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['time_h'] + [f'r={r*100:g}cm' for r in body_r])
        writer.writerows([[f'{t/3600:.4f}']+[f'{c:.4f}' for c in row] for t, row in zip(body_t, body_c)])
    text = '| 时间/h | ' + ' | '.join(f'{r*100:g} cm' for r in body_r) + ' |\n'
    text += '|---|---|---|---|---|---|\n'
    text += '\n'.join('| '+' | '.join([f'{t/3600:.4f}']+[f'{c:.4f}' for c in row])+' |' for t, row in zip(body_t, body_c))
    text += '\n\n含水率单位kg/kg。末行为临界状态，严格低于由事件后的真实积分确认。\n'
    text += '本次求解成功与导出校验不替代空间/时间收敛或物理验证。\n'
    (tables / 'table5.md').write_text(text, encoding='utf-8')
    np.savez_compressed(tables / 'table5_unrounded.npz', time_s=body_t, radius_m=body_r, moisture=body_c)
    check = dict(status='PASS', scope='export_readback_only', run_id=run_id,
                 rows=len(times), radii=21, table5_rows=len(body_t),
                 stored_value_max_error=float(np.max(abs(stored[:, 1:] - values))),
                 last_event_time_s=end, final_nonregular_row=abs(end/60-round(end/60)) > 1e-6/60,
                 storage='unrounded numeric; 0.0000 display', sha256=sha(path))
    save_json(folder / 'export_check.json', check)
    print(f'临界时间：{end/3600:.8f} h；结果已写入 {path}', flush=True)
    return check


# 9. 命令行入口。导入本文件不会求解；--check-only 也不会创建运行结果。
def main(argv=None):
    parser = argparse.ArgumentParser(description='第三问核心求解：单文件 + 附件1.xlsx')
    parser.add_argument('--input', type=Path, help='原始附件1.xlsx')
    parser.add_argument('--output', type=Path, help='输出根目录，默认本文件旁的单文件结果')
    parser.add_argument('--config', type=Path, help='Config同名字段的完整或部分JSON配置')
    parser.add_argument('--run-id', default='baseline', help='运行文件夹名称；改参数时使用新名称')
    parser.add_argument('--template', type=Path, help='可选的题目result3.xlsx模板')
    parser.add_argument('--resume', type=Path, help='本单文件产生的检查点npz')
    parser.add_argument('--branch', action='store_true', help='恢复时允许改变所支持的后段时间设置')
    modes = parser.add_mutually_exclusive_group()
    modes.add_argument('--check-only', action='store_true', help='只检查依赖、配置、附件，不积分')
    modes.add_argument('--export-only', action='store_true', help='只从已有完成运行导表')
    parser.add_argument('--N', type=int)
    parser.add_argument('--method', choices=('BDF', 'Radau'))
    parser.add_argument('--scenario', choices=SCENARIOS)
    for field in ('end', 'h', 'hm', 'rtol', 'atol_T', 'atol_C', 'max_step',
                  'early_rtol', 'early_atol_T', 'early_atol_C', 'early_max_step'):
        parser.add_argument('--'+field.replace('_', '-'), dest=field, type=float)
    args = parser.parse_args(argv)
    checked_run_id(args.run_id)
    if args.export_only:
        if args.resume or args.config or args.branch or any(getattr(args, field, None) is not None for field in Config.__dataclass_fields__):
            parser.error('--export-only使用原运行配置，不能同时更改配置或恢复。')
        export_run(args.run_id, args.output, args.template)
        return
    settings = asdict(Config())
    if args.resume:
        _, _, _, meta = read_checkpoint(args.resume)
        settings.update(meta['config'])
    if args.config:
        settings.update(json.loads(args.config.read_text('utf-8-sig')))
    for field in Config.__dataclass_fields__:
        value = getattr(args, field, None)
        if value is not None:
            settings[field] = value
    p = Config(**settings).check()
    source = find_input(args.input)
    env = read_environment(p.scenario, source)
    if args.template:
        if not args.template.is_file():
            raise FileNotFoundError(f'模板不存在：{args.template}')
        wb = load_workbook(args.template, read_only=True)
        try:
            if wb.sheetnames != ['Sheet1']:
                raise ValueError('模板须仅含Sheet1。')
        finally:
            wb.close()
    if args.branch and not args.resume:
        parser.error('--branch必须与--resume一起使用。')
    if args.check_only:
        print(json.dumps(dict(status='PASS',scope='config_and_input_only_no_integration',
            input=str(source),input_sha256=sha(source),observation_points=len(env.time),
            plateau_T_C=env.plateau_T,plateau_C=env.plateau_C,config=asdict(p),
            output=str(output_path(args.output))),ensure_ascii=False,indent=2))
        return
    result = run(p, args.run_id, resume=args.resume, branch=args.branch,
                 input_path=source, output_root=args.output)
    if result['event_found']:
        export_run(args.run_id, args.output, args.template)
    else:
        print(f'到{result["end_s"]/3600:g}h尚未找到临界事件，原始场与检查点已保存。')
        print('可增大end并从latest.npz恢复；请使用新的run-id。')


if __name__ == '__main__':
    main()

