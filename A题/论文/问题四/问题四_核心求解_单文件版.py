# -*- coding: utf-8 -*-
"""2026 A题第四问：移动材料坐标热湿模型，单文件核心。

本文件集中包含配置、输入、物性、几何、有限体积、推进、事件、恢复和导出。
求解所需数据为附件1.xlsx与附件2.xlsx；题目result4.xlsx模板可选。
内部温度为摄氏度、时间为秒、半径为米、含水率为kg/kg。
执行方案中的验证、制图、报告由旁侧脚本调用本核心，生产求解不依赖它们。
"""
from __future__ import annotations

# === TEAM PARAMETERS BEGIN ===
# 队友修改参数请优先改这里；时间s、长度m、温度℃、含水率kg/kg。
TEAM_PARAMETERS = {'purpose': 'production',
 'model': 'q4_appendix4_material',
 'state_layout': 'blocked_T_then_C',
 'N': 5120,
 'grid_kind': 'uniform_material',
 'R0_m': 0.02,
 'L_m': 0.25,
 'T0_C': 28.0,
 'C0': 2.55,
 'h': 25.0,
 'hm': 8e-07,
 'radius_mode': 'pchip',
 'radius_amplitude': 1.0,
 'radius_tail': 'plateau',
 'radius_tail_fraction': 0.0,
 'radius_tail_tau_s': 86400.0,
 'environment': 'S60',
 'temperature_shift_C': 0.0,
 'moisture_shift': 0.0,
 'method': 'BDF',
 'rtol': 1e-09,
 'atol_T': 1e-10,
 'atol_C': 1e-12,
 'max_step_s': 300.0,
 'early_max_step_s': 10.0,
 'gauss_order': 5,
 'threshold': 0.15,
 'root_xtol_s': 1e-05,
 'event_probe_count': 5,
 'event_fine_gap_s': 1.0,
 'sample_s': 60.0,
 'checkpoint_s': 21600.0,
 'end_s': 259200.0,
 'stop_at_event': True,
 'extension_s': 86400.0,
 'post_initial_s': 120.0,
 'balance_probe_count': 17}
# === TEAM PARAMETERS END ===

# === Q4 CORE BEGIN ===
import argparse
import csv
import hashlib
import json
import os
import platform
import shutil
import sys
from collections import deque
from dataclasses import asdict, dataclass, replace
from datetime import datetime, timezone
from functools import lru_cache
from pathlib import Path
from time import perf_counter

for _thread_name in ('OMP_NUM_THREADS', 'OPENBLAS_NUM_THREADS', 'MKL_NUM_THREADS', 'NUMEXPR_NUM_THREADS'):
    os.environ[_thread_name] = '1'

import numpy as np
import scipy
from numpy.polynomial.legendre import leggauss
from openpyxl import Workbook, load_workbook
from scipy.integrate import BDF, Radau, simpson
from scipy.interpolate import PchipInterpolator
from scipy.optimize import brentq
from scipy.sparse import bmat, diags

ROOT = Path(__file__).resolve().parent
CORE_PROTOCOL = 'q4_material_blocked_v1'


def now():
    return datetime.now(timezone.utc).isoformat()


def peak_memory_mb():
    if os.name == 'nt':
        import ctypes
        from ctypes import wintypes
        class Counters(ctypes.Structure):
            _fields_=[('cb',wintypes.DWORD),('PageFaultCount',wintypes.DWORD),
                *[(name,ctypes.c_size_t) for name in ('PeakWorkingSetSize','WorkingSetSize',
                'QuotaPeakPagedPoolUsage','QuotaPagedPoolUsage','QuotaPeakNonPagedPoolUsage',
                'QuotaNonPagedPoolUsage','PagefileUsage','PeakPagefileUsage')]]
        memory=Counters();memory.cb=ctypes.sizeof(memory)
        handle=ctypes.windll.kernel32.GetCurrentProcess
        handle.restype=wintypes.HANDLE
        get_memory=ctypes.windll.psapi.GetProcessMemoryInfo
        get_memory.argtypes=[wintypes.HANDLE,ctypes.POINTER(Counters),wintypes.DWORD]
        get_memory.restype=wintypes.BOOL
        if get_memory(handle(),ctypes.byref(memory),memory.cb):
            return float(memory.PeakWorkingSetSize/1024**2)
        return None
    import resource
    return float(resource.getrusage(resource.RUSAGE_SELF).ru_maxrss/(1024**2 if sys.platform=='darwin' else 1024))


def file_hash(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda: f.read(8 * 1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def json_hash(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                    separators=(',', ':'), allow_nan=False).encode('utf-8')).hexdigest()


def core_hash():
    """核心区域逐字节哈希；导出/CLI/图注在区域外，另记整文件哈希。"""
    lines = Path(__file__).read_text('utf-8').splitlines(keepends=True)
    start = next(i for i, x in enumerate(lines) if x.strip() == '# === Q4 CORE BEGIN ===')
    end = next(i for i, x in enumerate(lines) if x.strip() == '# === Q4 CORE END ===')
    return hashlib.sha256(''.join(lines[start:end + 1]).encode('utf-8')).hexdigest()


def save_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + '.tmp')
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
    temporary.replace(path)


def atomic_npz(path, **arrays):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.stem + '.tmp.npz')
    np.savez_compressed(temporary, **arrays)
    temporary.replace(path)


def read_json(path):
    return json.loads(Path(path).read_text('utf-8-sig'))


def update_stage(stage, status, evidence=None, next_action=None, root=ROOT, **extra):
    path = Path(root) / 'execution_status.json'
    state = read_json(path) if path.exists() else dict(
        plan_version='q4_execution_v1.0_single_core', current_stage='P00',
        stages={f'P{i:02}': 'pending' for i in range(12)}, latest_evidence=[],
        next_action='', latest_valid_checkpoint=None, active_run_id=None,
        failure_reason=None, pending_issues=[], release_status='NOT_RUN')
    state['current_stage'] = stage
    state['stages'][stage] = status
    if evidence is not None:
        state['latest_evidence'] = evidence
    if next_action is not None:
        state['next_action'] = next_action
    state.update(extra)
    state['updated_utc'] = now()
    save_json(path, state)


@dataclass(frozen=True)
class Config:
    # 这些是方案的探索起点；正式参数由通过验证的selected记录锁定。
    purpose: str = 'production'
    model: str = 'q4_appendix4_material'
    state_layout: str = 'blocked_T_then_C'
    N: int = 320
    grid_kind: str = 'uniform_material'
    R0_m: float = .02
    L_m: float = .25
    T0_C: float = 28.
    C0: float = 2.55
    h: float = 25.
    hm: float = 8e-7
    radius_mode: str = 'pchip'
    radius_amplitude: float = 1.
    radius_tail: str = 'plateau'
    radius_tail_fraction: float = 0.
    radius_tail_tau_s: float = 86400.
    environment: str = 'S60'
    temperature_shift_C: float = 0.
    moisture_shift: float = 0.
    method: str = 'BDF'
    rtol: float = 1e-6
    atol_T: float = 1e-7
    atol_C: float = 1e-9
    max_step_s: float = 300.
    early_max_step_s: float = 10.
    gauss_order: int = 5
    threshold: float = .15
    root_xtol_s: float = 1e-5
    event_probe_count: int = 5
    event_fine_gap_s: float = 1.
    sample_s: float = 60.
    checkpoint_s: float = 21600.
    end_s: float = 259200.
    stop_at_event: bool = True
    extension_s: float = 86400.
    post_initial_s: float = 10.
    balance_probe_count: int = 17

    def validate(self):
        if self.purpose not in ('production', 'verification'):
            raise ValueError('purpose只能为production或verification')
        if (self.model, self.state_layout, self.grid_kind) != (
                'q4_appendix4_material', 'blocked_T_then_C', 'uniform_material'):
            raise ValueError('模型/状态布局/材料网格标记不匹配')
        integers = dict(N=2, gauss_order=1, event_probe_count=5, balance_probe_count=17)
        for name, minimum in integers.items():
            v = getattr(self, name)
            if isinstance(v, bool) or not isinstance(v, int) or v < minimum:
                raise ValueError(f'{name}须为至少{minimum}的整数')
        if self.balance_probe_count not in (17, 33, 65):
            raise ValueError('平衡探针数使用17、33或65，以支持隔点Simpson细化')
        if not isinstance(self.stop_at_event, bool):
            raise ValueError('stop_at_event必须为布尔量')
        for name, value in asdict(self).items():
            if isinstance(value, (int, float)) and not np.isfinite(value):
                raise ValueError(f'非有限配置：{name}')
        positive = ('R0_m', 'L_m', 'C0', 'radius_amplitude', 'radius_tail_tau_s',
                    'rtol', 'atol_T', 'atol_C', 'max_step_s', 'early_max_step_s',
                    'threshold', 'root_xtol_s', 'event_fine_gap_s', 'sample_s',
                    'checkpoint_s', 'end_s', 'extension_s', 'post_initial_s')
        if any(getattr(self, x) <= 0 for x in positive):
            raise ValueError('尺寸、容差、时间设置和含水率必须为正')
        if self.h < 0 or self.hm < 0 or (self.purpose == 'production' and min(self.h, self.hm) <= 0):
            raise ValueError('生产h和hm须为正；零通量只用于verification')
        if self.T0_C <= -273.15 or self.C0 <= self.threshold:
            raise ValueError('初始温度或下降阈值配置不合法')
        if self.radius_mode not in ('pchip', 'linear', 'fixed') or self.radius_tail not in ('plateau', 'exp_shrink'):
            raise ValueError('半径模式/尾段模式不支持')
        if not 0 <= self.radius_tail_fraction < 1:
            raise ValueError('尾段收缩比例须位于[0,1)')
        if self.radius_tail == 'plateau' and self.radius_tail_fraction != 0:
            raise ValueError('plateau不使用非零尾段收缩比例')
        if self.radius_mode == 'fixed' and (self.radius_amplitude != 1 or self.radius_tail != 'plateau'):
            raise ValueError('固定半径对照不能同时改变收缩幅度或尾段')
        if self.environment not in ('S0', 'S30', 'S60', 'S90') or self.method not in ('BDF', 'Radau'):
            raise ValueError('环境或求解器名称不支持')
        if self.sample_s > 60 or not np.isclose(60 / self.sample_s, round(60 / self.sample_s), rtol=0, atol=1e-10):
            raise ValueError('全场保存间隔必须整除60s，避免规定输出时刻缺失')
        if self.post_initial_s < 10:
            raise ValueError('须至少真实推进根后10s')
        return self

    @classmethod
    def read(cls, path):
        return cls(**read_json(path)).validate()


def stable_copy(source, destination):
    source, destination = Path(source).resolve(), Path(destination).resolve()
    before = file_hash(source)
    if destination.exists():
        if file_hash(destination) != before:
            raise ValueError(f'已有冻结副本不同，保留双方文件：{destination}')
    else:
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, destination)
    after, copied = file_hash(source), file_hash(destination)
    if not before == after == copied:
        raise ValueError(f'复制期间源发生变化：{source}')
    return dict(source=str(source), copy=str(destination), bytes=source.stat().st_size,
                source_sha256=before, copy_sha256=copied, copied_utc=now())


def find_input(name, root=ROOT, explicit=None):
    if explicit is not None:
        path = Path(explicit).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(path)
        return path
    base = Path(root).resolve()
    for path in (base / 'inputs' / name, base / name, base / '附件' / name, base.parent / '附件' / name):
        if path.is_file():
            return path
    raise FileNotFoundError(f'缺少{name}，请放到核心文件旁或inputs目录')


def read_excel_numeric(path, columns, count, last_time, interval):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb.active.values)
    finally:
        wb.close()
    body = [row[:columns] for row in rows[1:] if any(x is not None for x in row)]
    a = np.asarray(body, dtype=np.float64)
    if a.shape != (count, columns) or not np.isfinite(a).all():
        raise ValueError(f'{path}数据形状/有限值不符合约定：{a.shape}')
    if a[0, 0] != 0 or a[-1, 0] != last_time or not np.array_equal(np.diff(a[:, 0]), np.full(count - 1, interval)):
        raise ValueError(f'{path}时间必须从0至{last_time}，严格按{interval}s递增')
    return a


def prepare_inputs(project_root=ROOT.parent, q4_root=ROOT):
    project_root, q4_root = Path(project_root), Path(q4_root)
    manifest_path = q4_root / 'inputs/input_manifest.json'
    if manifest_path.exists():
        old = read_json(manifest_path)
        for item in old['files']:
            if file_hash(q4_root / item['copy']) != item['copy_sha256']:
                raise ValueError('冻结输入副本哈希改变')
            if Path(item['source']).is_file() and file_hash(item['source']) != item['source_sha256']:
                raise ValueError('原始来源已改变，请核对新版本后使用新快照')
        if (q4_root/'reference/q3_snapshot/snapshot_manifest.json').exists() and (q4_root/'configs/candidate.json').exists():
            return old
    manual = q4_root / 'A题问题四_建模编程论文交接手册_全面修订版_v2.0.docx'
    pairs = [(project_root/'A题.pdf', 'inputs/A题.pdf'),
             (project_root/'附件/附件1.xlsx', 'inputs/附件1.xlsx'),
             (project_root/'附件/附件2.xlsx', 'inputs/附件2.xlsx'),
             (project_root/'附件/附件3/result4.xlsx', 'inputs/result4_template.xlsx'),
             (manual, 'inputs/建模依据_v2.0.docx')]
    files = []
    for source, copy in pairs:
        item = stable_copy(source, q4_root / copy)
        item['copy'] = copy
        files.append(item)
    env = read_excel_numeric(q4_root/'inputs/附件1.xlsx', 3, 241, 14400, 60)
    radius = read_excel_numeric(q4_root/'inputs/附件2.xlsx', 2, 145, 259200, 1800)
    result = dict(schema='q4_input_v1', files=files,
                  environment=dict(rows=len(env), time_unit='s', T_unit='degC', C_unit='kg/kg', time_range_s=[0,14400]),
                  radius=dict(rows=len(radius), time_unit='s', original_radius_unit='cm', internal_radius_unit='m',
                              time_range_s=[0,259200], initial_m=float(radius[0,1]/100), final_m=float(radius[-1,1]/100)))
    references = list((project_root/'问题三/q3').glob('*.py'))
    references += [project_root/'问题三'/x for x in (
        'configs/accepted_config.json','outputs/q3/formal_manifest.json',
        'outputs/q3/validation/regression.json','outputs/q3/validation/manufactured.json',
        'outputs/q3/validation/convergence.json','outputs/q3/validation/scenario_validation.json')]
    ref_records = []
    for source in references:
        copy = Path('reference/q3_snapshot') / source.relative_to(project_root/'问题三')
        item = stable_copy(source, q4_root/copy)
        item['copy'] = str(copy)
        ref_records.append(item)
    save_json(q4_root/'reference/q3_snapshot/snapshot_manifest.json', dict(files=ref_records))
    save_json(q4_root/'configs/candidate.json', asdict(Config()))
    save_json(manifest_path, result)
    return result


def window_mean(t, values, width_s):
    left = t[-1] - width_s
    if width_s <= 0 or left < t[0]:
        raise ValueError('时间平均窗口超出观测范围')
    x = np.r_[left, t[(t > left) & (t < t[-1])], t[-1]]
    y = np.interp(x, t, values)
    return float(np.sum(np.diff(x) * (y[:-1] + y[1:]) / 2) / width_s)


class EnvironmentInput:
    def __init__(self, cfg, path):
        a = read_excel_numeric(path, 3, 241, 14400, 60)
        self.time_s, self.temperature_C, self.moisture = a.T
        if np.any(self.temperature_C <= -273.15) or np.any(self.moisture <= 0):
            raise ValueError('环境超出温湿合法范围')
        if cfg.environment == 'S0':
            T, C = self.temperature_C[-1], self.moisture[-1]
        else:
            width = {'S30':1800, 'S60':3600, 'S90':5400}[cfg.environment]
            T = window_mean(self.time_s, self.temperature_C, width)
            C = window_mean(self.time_s, self.moisture, width)
        self.plateau_T = float(T + cfg.temperature_shift_C)
        self.plateau_C = float(C + cfg.moisture_shift)
        if self.plateau_T <= -273.15 or self.plateau_C <= 0:
            raise ValueError('环境偏移导致平台非法')

    def at(self, t_s, side):
        t = np.asarray(t_s, dtype=float)
        if not np.isfinite(t).all() or np.any(t < 0):
            raise ValueError('环境时间非法')
        if side == 'observed':
            if np.any(t > self.time_s[-1]):
                raise ValueError('观测侧禁止越过4h')
            return np.interp(t, self.time_s, self.temperature_C), np.interp(t, self.time_s, self.moisture)
        if side != 'plateau' or np.any(t < self.time_s[-1]):
            raise ValueError('平台侧只能在4h及以后使用')
        return np.full_like(t, self.plateau_T), np.full_like(t, self.plateau_C)


class RadiusInput:
    def __init__(self, cfg, path):
        self.cfg = cfg
        a = read_excel_numeric(path, 2, 145, 259200, 1800)
        self.time_s, radius_cm = a.T
        self.radius_m = radius_cm / 100
        if np.any(self.radius_m <= 0) or np.any(np.diff(self.radius_m) > 0) or abs(self.radius_m[0] - cfg.R0_m) > 1e-12:
            raise ValueError('半径必须为正、单调不增，且起点与R0一致')
        self.curve = PchipInterpolator(self.time_s, self.radius_m, extrapolate=False)
        self.last_s = float(self.time_s[-1])
        self.tail_R = cfg.R0_m - cfg.radius_amplitude * (cfg.R0_m - self.radius_m[-1])
        if self.tail_R <= 0:
            raise ValueError('幅度调整导致半径非正')

    def at(self, t_s):
        t = np.asarray(t_s, dtype=float)
        if not np.isfinite(t).all() or np.any(t < 0):
            raise ValueError('半径时间非法')
        p = self.cfg
        if p.radius_mode == 'fixed':
            return np.full_like(t, p.R0_m)
        within = np.minimum(t, self.last_s)
        base = self.curve(within) if p.radius_mode == 'pchip' else np.interp(within, self.time_s, self.radius_m)
        result = p.R0_m - p.radius_amplitude * (p.R0_m - base)
        if p.radius_tail == 'exp_shrink':
            tail = self.tail_R * (1 + p.radius_tail_fraction * np.expm1(-np.maximum(t-self.last_s,0)/p.radius_tail_tau_s))
            result = np.where(t > self.last_s, tail, result)
        if np.any(result <= 0):
            raise ValueError('半径非正')
        return result

    def derivative(self, t_s, side='right'):
        t = np.asarray(t_s, dtype=float)
        self.at(t)
        if side not in ('left', 'right'):
            raise ValueError('导数侧须为left/right')
        p = self.cfg
        if p.radius_mode == 'fixed':
            return np.zeros_like(t)
        within = np.minimum(t, self.last_s)
        if p.radius_mode == 'pchip':
            derivative = self.curve.derivative()(within)
        else:
            i = np.clip(np.searchsorted(self.time_s, within, side=side)-1, 0, len(self.time_s)-2)
            derivative = np.diff(self.radius_m)[i] / np.diff(self.time_s)[i]
        derivative = derivative * p.radius_amplitude
        tail = np.zeros_like(t)
        if p.radius_tail == 'exp_shrink':
            tail = -self.tail_R*p.radius_tail_fraction/p.radius_tail_tau_s*np.exp(-np.maximum(t-self.last_s,0)/p.radius_tail_tau_s)
        outside = t >= self.last_s if side == 'right' else t > self.last_s
        return np.where(outside, tail, derivative)


def rho(C):
    return 760. + 90.*C


def cp(C):
    return 1850. + 2150.*C/(C+1.)


def k(C):
    return .12 + .20*C/(C+1.)


def B(C):
    return rho(C)*cp(C)


def A(T_K):
    return 4.2e-4*np.exp(-3850./T_K)


def D(C,T_K):
    return A(T_K)*np.exp(-.30/C)


def k_C(C):
    return .20/(C+1.)**2


def D_C(C,T_K):
    return .30*D(C,T_K)/C**2


def D_T(C,T_K):
    return 3850.*D(C,T_K)/T_K**2


@dataclass(frozen=True)
class MaterialGrid:
    xi: np.ndarray
    xi_faces: np.ndarray


@dataclass(frozen=True)
class Geometry:
    R_m: float
    r_m: np.ndarray
    faces_m: np.ndarray
    areas_m2: np.ndarray
    volumes_m3: np.ndarray
    J: float


def make_grid(N):
    xi = np.linspace(0.,1.,N+1)
    faces = np.r_[0.,(xi[:-1]+xi[1:])/2,1.]
    xi.flags.writeable = False
    faces.flags.writeable = False
    return MaterialGrid(xi,faces)


def geometry(grid, R_m, R0_m, L_m):
    R_m = float(R_m)
    if R_m <= 0 or not np.isfinite(R_m):
        raise ValueError('几何半径非法')
    return Geometry(R_m, R_m*grid.xi, R_m*grid.xi_faces,
                    2*np.pi*L_m*R_m*grid.xi_faces,
                    np.pi*L_m*R_m**2*np.diff(grid.xi_faces**2), (R_m/R0_m)**2)


def pack(T, C):
    T,C = np.asarray(T),np.asarray(C)
    if T.ndim != 1 or T.shape != C.shape:
        raise ValueError('T和C须为一维等长数组')
    return np.concatenate((T,C)).astype(np.float64,copy=False)


def unpack(y, n):
    y = np.asarray(y)
    if y.ndim != 1 or y.size != 2*n:
        raise ValueError('分块状态长度不匹配')
    return y[:n],y[n:]


def make_atol(n, atol_T, atol_C):
    return np.r_[np.full(n,atol_T),np.full(n,atol_C)]


def sparsity(n):
    G = diags([np.ones(n-1),np.ones(n),np.ones(n-1)],[-1,0,1],format='csc')
    return bmat([[G,G],[G,G]],format='csc').astype(bool)


def reconstruct_physical(C, xi, R_m, targets_m):
    target = np.asarray(targets_m,dtype=float)
    if R_m <= 0 or np.any(target < 0) or not np.isfinite(target).all():
        raise ValueError('物理位置/半径非法')
    valid = target <= R_m+1e-12
    values = np.full(target.shape,np.nan)
    values[valid] = np.interp(np.minimum(target[valid],R_m)/R_m,xi,C)
    return values,valid


@lru_cache(None)
def gauss_rule(order):
    points,weights = leggauss(order)
    return (points+1)/2,weights/2


def delta_phi(left, right, order=5):
    left,right = np.asarray(left),np.asarray(right)
    delta = right-left
    s,w = gauss_rule(order)
    return delta*np.sum(w[:,None]*np.exp(-.30/(left[None,:]+s[:,None]*delta[None,:])),axis=0)


def fluxes(T,C,grid,geom,cfg,env_values):
    ft,fc = np.zeros(len(T)+1),np.zeros(len(C)+1)
    kval = k(C)
    kface = 2*kval[:-1]*kval[1:]/(kval[:-1]+kval[1:])
    factor = 2*np.pi*cfg.L_m*grid.xi_faces[1:-1]/np.diff(grid.xi)
    ft[1:-1] = -factor*kface*np.diff(T)
    fc[1:-1] = -factor*A((T[:-1]+T[1:])/2+273.15)*delta_phi(C[:-1],C[1:],cfg.gauss_order)
    ft[-1] = geom.areas_m2[-1]*cfg.h*(T[-1]-env_values[0])
    fc[-1] = geom.areas_m2[-1]*cfg.hm*(C[-1]-env_values[1])
    return ft,fc


def legal_state(y,n,t):
    T,C = unpack(y,n)
    if not np.isfinite(y).all() or np.any(C <= 0) or np.any(T <= -273.15):
        raise ValueError(f'非法温湿状态t={t}, minC={np.min(C)}, minT={np.min(T)}')
    return T,C


def make_rhs(cfg,grid,radius,environment,environment_side,source=None):
    if source is not None and cfg.purpose != 'verification':
        raise ValueError('生产方程不得含制造源项')
    n = len(grid.xi)
    def rhs(t,y):
        T,C = legal_state(y,n,t)
        geom = geometry(grid,radius.at(t),cfg.R0_m,cfg.L_m)
        ft,fc = fluxes(T,C,grid,geom,cfg,environment.at(t,environment_side))
        net_T,net_C = -np.diff(ft),-np.diff(fc)
        if source is not None:
            ST,SC = source(t,grid.xi)
            net_T += geom.volumes_m3*ST
            net_C += geom.volumes_m3*SC
        return pack(net_T/(B(C)*geom.volumes_m3),net_C/geom.volumes_m3)
    return rhs

@dataclass
class StepRecord:
    t0: float
    t1: float
    y0: np.ndarray
    y1: np.ndarray
    interpolant: object
    nfev: int
    njev: int
    nlu: int

    def dense(self, t):
        a = np.asarray(t)
        if np.any(a < self.t0-1e-9) or np.any(a > self.t1+1e-9):
            raise ValueError('禁止在接受步区间之外使用稠密输出')
        return self.interpolant(np.clip(a,self.t0,self.t1))


def advance_segment(rhs, t0, y0, t1, cfg, max_step=None):
    """每次只保留一个接受步的稠密输出；向外提供真实的接受步。"""
    if t1 <= t0:
        return
    method = {'BDF':BDF,'Radau':Radau}[cfg.method]
    solver = method(rhs,float(t0),np.asarray(y0,dtype=float),float(t1),
        rtol=cfg.rtol,atol=make_atol(cfg.N+1,cfg.atol_T,cfg.atol_C),
        max_step=cfg.max_step_s if max_step is None else max_step,
        first_step=min(.001,float(t1-t0),cfg.max_step_s if max_step is None else max_step),
        jac_sparsity=sparsity(cfg.N+1),vectorized=False)
    stats = (0,0,0)
    while solver.status == 'running':
        a,ya = float(solver.t),solver.y.copy()
        message = solver.step()
        if solver.status == 'failed':
            raise RuntimeError(f'隐式积分失败t={a}: {message}')
        b,yb = float(solver.t),solver.y.copy()
        legal_state(yb,cfg.N+1,b)
        current = (solver.nfev,solver.njev,solver.nlu)
        record = StepRecord(a,b,ya,yb,solver.dense_output(),*(x-y for x,y in zip(current,stats)))
        stats = current
        yield record


def segment_ends(t0,end,cfg):
    """环境折点、4h两侧、线性半径折点、72h和检查点均显式分段。"""
    candidates = [end,14400.,259200.]
    candidates.extend(np.arange(60.,14400.1,60.))
    candidates.extend(np.arange(cfg.checkpoint_s,end+cfg.checkpoint_s,cfg.checkpoint_s))
    if cfg.radius_mode == 'linear':
        candidates.extend(np.arange(1800.,259200.1,1800.))
    return sorted({float(t) for t in candidates if t0+1e-8 < t <= end})


def max_info(y,grid,R):
    C = unpack(y,len(grid.xi))[1]
    i = int(np.argmax(C))
    return dict(Cmax=float(C[i]),node=i,xi=float(grid.xi[i]),r_m=float(R*grid.xi[i]),
                center_is_max=bool(C[0] >= C[i]-1e-12))


def scan_event(step,cfg):
    """全部节点的最大值：先5探针，近阈值时相邻探针不超过1s。"""
    n = cfg.N+1
    ts = np.linspace(step.t0,step.t1,cfg.event_probe_count)
    gs = np.max(step.dense(ts)[n:,:],axis=0)-cfg.threshold
    if np.min(gs) <= .02 or np.any(gs[:-1]*gs[1:] <= 0):
        ts = np.unique(np.r_[ts,np.linspace(step.t0,step.t1,
                            max(2,int(np.ceil((step.t1-step.t0)/cfg.event_fine_gap_s))+1))])
        gs = np.max(step.dense(ts)[n:,:],axis=0)-cfg.threshold
    brackets = [(float(ts[i]),float(ts[i+1])) for i in range(len(ts)-1)
                if gs[i] >= 0 and gs[i+1] < 0]
    if not brackets:
        return None
    a,b = brackets[0]
    fun = lambda t: float(np.max(step.dense(t)[n:])-cfg.threshold)
    root = float(brentq(fun,a,b,xtol=cfg.root_xtol_s,rtol=4*np.finfo(float).eps))
    left,right = a,b
    while right-left > min(.6,max(cfg.root_xtol_s*2,1e-6)):
        middle = (left+right)/2
        if fun(middle) >= 0:
            left = middle
        else:
            right = middle
    return dict(time_s=root,time_h=root/3600,original_bracket_s=[a,b],
        bracket_s=[left,right],bracket_g=[fun(left),fun(right)],bracket_width_s=right-left,
        root_g=fun(root),candidate_brackets_s=brackets,probe_count=len(ts),
        strict_post_pass=None,epsilon_C=None)


class RunRecorder:
    def __init__(self,cfg,grid,radius,env,directory,resume=False):
        self.cfg,self.grid,self.radius,self.env,self.directory = cfg,grid,radius,env,Path(directory)
        self.times,self.states,self.radii,self.kinds = [],[],[],[]
        self.probes,self.steps = [],[]
        self.raw_index=[];self.probe_index=[]
        self.committed_rows=0;self.committed_probes=0
        self.step_id = 0
        self.interval_id = 0
        if resume:
            with np.load(self.directory/'fields.npz') as a:
                self.times = a['time_s'].tolist()
                self.states = [pack(T,C) for T,C in zip(a['T_C'],a['C'])]
                self.radii = a['R_m'].tolist()
                self.kinds = a['kind'].tolist()
            with np.load(self.directory/'balance_probes.npz') as a:
                self.probes = [x for x in a['probes']]
            with np.load(self.directory/'steps.npz') as a:
                self.steps = a['steps'].tolist()
            if self.probes:
                self.interval_id = int(self.probes[-1][0,5])+1
            self.step_id = len(self.steps)
            self.raw_index=read_json(self.directory/'raw/index.json')['chunks']
            self.probe_index=read_json(self.directory/'balance_probe/index.json')['chunks']
            self.committed_rows=len(self.times);self.committed_probes=len(self.probes)

    def add(self,t,y,kind):
        legal_state(y,self.cfg.N+1,t)
        if self.times and abs(self.times[-1]-t) <= 1e-7:
            if kind == 'event':
                self.times[-1],self.states[-1],self.radii[-1],self.kinds[-1] = float(t),y.copy(),float(self.radius.at(t)),kind
            return
        if self.times and t < self.times[-1]:
            raise ValueError('保存时间顺序错误')
        self.times.append(float(t)); self.states.append(y.copy())
        self.radii.append(float(self.radius.at(t))); self.kinds.append(kind)

    def accept(self,step,until,side,extra=()):
        cfg = self.cfg
        regular = np.arange((np.floor(step.t0/cfg.sample_s)+1)*cfg.sample_s,
                            until+1e-8,cfg.sample_s)
        endpoints = [(float(t),'regular') for t in regular if t <= until]
        endpoints.extend((float(t),kind) for t,kind in extra if step.t0 < t <= until+1e-8)
        endpoints.sort(key=lambda x:x[0])
        unique = {}
        for t,kind in endpoints:
            unique[t] = kind
        split = sorted({step.t0,until,*unique})
        for a,b in zip(split[:-1],split[1:]):
            if b <= a:
                continue
            ts = np.linspace(a,b,cfg.balance_probe_count)
            cs = step.dense(ts)[-1,:]
            rr = self.radius.at(ts)
            ce = self.env.at(ts,side)[1]
            self.probes.append(np.column_stack([ts,cs,rr,ce,
                np.full_like(ts,self.step_id),np.full_like(ts,self.interval_id),
                np.full_like(ts,0 if side=='observed' else 1)]))
            self.interval_id += 1
        for t,kind in unique.items():
            self.add(t,step.dense(t),kind)
        state=step.dense(until);T,C=unpack(state,cfg.N+1);R=float(self.radius.at(until))
        geom=geometry(self.grid,R,cfg.R0_m,cfg.L_m);i=int(np.argmax(C))
        fs=2*np.pi*cfg.L_m*R*cfg.hm*(C[-1]-self.env.at(until,side)[1])
        self.steps.append([step.t0,until,step.t1,step.nfev,step.njev,step.nlu,
            0 if side=='observed' else 1,R,geom.J,C[i],C[0],C[-1],np.min(C),
            self.grid.xi[i],R*self.grid.xi[i],float(geom.volumes_m3@C),float(fs),np.min(T),np.max(T)])
        self.step_id += 1

    def flush(self,t,y,manifest,event):
        a = np.asarray(self.states,dtype=np.float64)
        n = self.cfg.N+1
        if len(self.times)>self.committed_rows:
            start=self.committed_rows;rel=f'raw/chunk_{len(self.raw_index):05}.npz'
            atomic_npz(self.directory/rel,time_s=np.asarray(self.times[start:]),xi=self.grid.xi,
                R_m=np.asarray(self.radii[start:]),T_C=a[start:,:n],C=a[start:,n:],kind=np.asarray(self.kinds[start:]))
            self.raw_index.append(dict(path=rel,sha256=file_hash(self.directory/rel),
                time_range_s=[self.times[start],self.times[-1]],shape=[len(self.times)-start,n]))
            self.committed_rows=len(self.times)
        if len(self.probes)>self.committed_probes:
            start=self.committed_probes;rel=f'balance_probe/chunk_{len(self.probe_index):05}.npz'
            array=np.asarray(self.probes[start:],dtype=float)
            atomic_npz(self.directory/rel,probes=array)
            self.probe_index.append(dict(path=rel,sha256=file_hash(self.directory/rel),
                time_range_s=[float(array[0,0,0]),float(array[-1,-1,0])],shape=list(array.shape)))
            self.committed_probes=len(self.probes)
        save_json(self.directory/'raw/index.json',dict(chunks=self.raw_index,
            units=dict(time_s='s',xi='1',R_m='m',T_C='degC',C='kg/kg'),rows=len(self.times)))
        save_json(self.directory/'balance_probe/index.json',dict(chunks=self.probe_index,
            columns=['time_s','C_surface','R_m','Ce','step_id','subinterval_id','environment_side']))
        atomic_npz(self.directory/'fields.npz',time_s=np.asarray(self.times),xi=self.grid.xi,
                   R_m=np.asarray(self.radii),T_C=a[:,:n],C=a[:,n:],kind=np.asarray(self.kinds))
        atomic_npz(self.directory/'balance_probes.npz',probes=np.asarray(self.probes,dtype=float).reshape(-1,self.cfg.balance_probe_count,7))
        atomic_npz(self.directory/'steps.npz',steps=np.asarray(self.steps,dtype=float).reshape(-1,19))
        with (self.directory/'diagnostics.csv').open('w',newline='',encoding='utf-8-sig') as f:
            w=csv.writer(f);w.writerow(['t0_s','t1_s','accepted_end_s','nfev','njev','nlu','environment_side',
                'R_m','J','Cmax','Ccenter','Csurface','Cmin','xi_max','r_max_m','I_C','F_surface','Tmin_C','Tmax_C'])
            w.writerows(self.steps)
        # 保存的是实际合法状态，而不是可能尚未接受的牛顿试探状态。
        atomic_npz(self.directory/'checkpoint.npz',time_s=np.array(t),state=y,xi=self.grid.xi)
        checkpoint = dict(time_s=float(t),state_layout=self.cfg.state_layout,config_sha256=json_hash(asdict(self.cfg)),
            core_sha256=manifest['core_sha256'],input_sha256=manifest['input_sha256'],run_id=manifest['run_id'],
            rows=len(self.times),steps=len(self.steps),event=event,updated_utc=now(),
            state_sha256=file_hash(self.directory/'checkpoint.npz'),
            history_sha256={name:file_hash(self.directory/name) for name in ('fields.npz','balance_probes.npz','steps.npz',
                'raw/index.json','balance_probe/index.json')})
        save_json(self.directory/'checkpoint.json',checkpoint)
        manifest.update(last_valid_time_s=float(t),last_valid_checkpoint='checkpoint.npz',
                        rows=len(self.times),accepted_steps=len(self.steps),updated_utc=now(),peak_memory_mb=peak_memory_mb())
        save_json(self.directory/'manifest.json',manifest)
        if event is not None:
            save_json(self.directory/'event.json',event)


def run_model(cfg,run_id,root=ROOT,resume=False,env_path=None,radius_path=None):
    cfg.validate()
    root = Path(root)
    if not run_id or any(x in run_id for x in ('/','\\','..')):
        raise ValueError('run_id须为单一目录名称')
    directory = root/'runs'/run_id
    env_path = find_input('附件1.xlsx',root,env_path)
    radius_path = find_input('附件2.xlsx',root,radius_path)
    inputs = dict(environment=file_hash(env_path),radius=file_hash(radius_path))
    inputs_hash = json_hash(inputs)
    grid = make_grid(cfg.N)
    radius,env = RadiusInput(cfg,radius_path),EnvironmentInput(cfg,env_path)
    manifest_path = directory/'manifest.json'
    started = perf_counter()
    event = None
    if manifest_path.exists():
        if not resume:
            raise FileExistsError(f'运行已存在，禁止覆盖：{run_id}')
        manifest = read_json(manifest_path)
        old_cfg = read_json(directory/'config.json')
        changed = {k for k,v in asdict(cfg).items() if old_cfg.get(k) != v}
        if changed-{'end_s'} or cfg.end_s < old_cfg['end_s']:
            raise ValueError(f'恢复只允许增加end_s，禁止变更：{changed}')
        if manifest['core_sha256'] != core_hash() or manifest['input_sha256'] != inputs_hash:
            raise ValueError('恢复的代码或输入哈希不一致')
        checkpoint = read_json(directory/'checkpoint.json')
        if checkpoint['run_id'] != run_id or checkpoint['state_layout'] != cfg.state_layout:
            raise ValueError('检查点身份错误')
        if (checkpoint['config_sha256'] != json_hash(old_cfg) or checkpoint['core_sha256'] != manifest['core_sha256']
            or checkpoint['input_sha256'] != inputs_hash or checkpoint['state_sha256'] != file_hash(directory/'checkpoint.npz')):
            raise ValueError('检查点配置、核心、输入或状态哈希不匹配')
        for name,expected in checkpoint['history_sha256'].items():
            if file_hash(directory/name) != expected:
                raise ValueError(f'检查点历史原始数据哈希不匹配：{name}')
        for index_name in ('raw/index.json','balance_probe/index.json'):
            for chunk in read_json(directory/index_name)['chunks']:
                if file_hash(directory/chunk['path'])!=chunk['sha256']:
                    raise ValueError('历史分块哈希不一致：'+chunk['path'])
        with np.load(directory/'checkpoint.npz') as a:
            t,y = float(a['time_s']),a['state'].copy()
            if not np.array_equal(a['xi'],grid.xi):
                raise ValueError('检查点材料网格不同')
        if t >= cfg.end_s:
            raise ValueError('恢复上限没有超过检查点')
        event = checkpoint['event']
        if event is not None:
            raise ValueError('已完成事件的运行无需延长；扩展后验实验请用新运行')
        manifest.setdefault('resume_history',[]).append(dict(config=old_cfg,config_sha256=json_hash(old_cfg),
            checkpoint_time_s=t,resumed_utc=now()))
    else:
        if resume:
            raise FileNotFoundError('没有可恢复运行')
        directory.mkdir(parents=True)
        t,y = 0.,pack(np.full(cfg.N+1,cfg.T0_C),np.full(cfg.N+1,cfg.C0))
        manifest = dict(schema='q4_run_v1',run_id=run_id,protocol=CORE_PROTOCOL,
            core_sha256=core_hash(),whole_file_sha256=file_hash(__file__),input_files=inputs,
            input_sha256=inputs_hash,created_utc=now(),runtime=dict(python=platform.python_version(),
            numpy=np.__version__,scipy=scipy.__version__,pid=os.getpid(),threads=1),
            environment_plateau=dict(T_C=env.plateau_T,C=env.plateau_C),resume_history=[])
    manifest.update(status='running',config_sha256=json_hash(asdict(cfg)),elapsed_this_call_s=0.)
    save_json(directory/'config.json',asdict(cfg))
    save_json(directory/'resolved_config.json',asdict(cfg))
    recorder = RunRecorder(cfg,grid,radius,env,directory,resume)
    if not resume:
        recorder.add(t,y,'initial')
    recorder.flush(t,y,manifest,event)
    recent = deque()
    post_targets = []
    neighborhood=[]
    end = cfg.end_s
    try:
        while t < end-1e-8:
            segment_end = segment_ends(t,end,cfg)[0]
            side = 'observed' if t < 14400 else 'plateau'
            rhs = make_rhs(cfg,grid,radius,env,side)
            max_step = cfg.early_max_step_s if t < 14400 else cfg.max_step_s
            if event is not None and cfg.stop_at_event:
                max_step = min(max_step,1.)
            found = False
            for step in advance_segment(rhs,t,y,segment_end,cfg,max_step):
                recent.append(step)
                while recent and recent[0].t1 < step.t0-30.:
                    recent.popleft()
                detected = scan_event(step,cfg) if event is None else None
                if detected is not None:
                    event = detected
                    root_t = event['time_s']
                    root_y = step.dense(root_t)
                    event['at_root'] = max_info(root_y,grid,float(radius.at(root_t)))
                    event['pre_checks'] = []
                    for dt in (10.,5.,2.,1.):
                        query = root_t-dt
                        for history in recent:
                            if history.t0 <= query <= history.t1:
                                event['pre_checks'].append(dict(time_s=query,offset_s=-dt,
                                    **max_info(history.dense(query),grid,float(radius.at(query)))))
                                neighborhood.append((query,history.dense(query),history.t0,history.t1,'pre'))
                                break
                    event['post_checks'] = []
                    neighborhood.append((root_t,root_y,step.t0,step.t1,'root'))
                    recorder.accept(step,root_t,side,[(root_t,'event')])
                    recorder.add(root_t,root_y,'event')
                    t,y = root_t,root_y
                    # 明确从根状态重新积分，不能把求根所在步向后外推当作后验。
                    if cfg.stop_at_event:
                        end = root_t+cfg.post_initial_s
                    post_targets = [(root_t+dt,'post') for dt in (1.,2.,5.,10.,20.,40.,80.,120.,160.,320.,600.)
                                    if root_t+dt <= end]
                    found = True
                    recorder.flush(t,y,manifest,event)
                    break
                extra = [(u,kind) for u,kind in post_targets if step.t0 < u <= step.t1]
                if abs(step.t1-segment_end) <= 1e-8:
                    extra.append((step.t1,'segment_end'))
                recorder.accept(step,step.t1,side,extra)
                for u,kind in extra:
                    if kind == 'post':
                        event['post_checks'].append(dict(time_s=u,offset_s=u-event['time_s'],
                            **max_info(step.dense(u),grid,float(radius.at(u)))))
                        neighborhood.append((u,step.dense(u),step.t0,step.t1,'post'))
                t,y = step.t1,step.y1
            if not found and (abs(t/cfg.checkpoint_s-round(t/cfg.checkpoint_s)) < 1e-10 or t >= end-1e-8):
                recorder.flush(t,y,manifest,event)
                print(f'{run_id}: {t/3600:.4f} h, Cmax={max_info(y,grid,float(radius.at(t)))["Cmax"]:.8f}',flush=True)
        manifest.update(status='completed',end_reason='event_with_post_integration' if event is not None else 'not_reached',
            elapsed_this_call_s=perf_counter()-started)
        if event is None:
            manifest['extension_assessment'] = dict(Cmax=max_info(y,grid,float(radius.at(t)))['Cmax'],
                environment_below_threshold=env.plateau_C < cfg.threshold,
                recent_Cmax_change=float(max_info(y,grid,float(radius.at(t)))['Cmax']-
                    np.max(recorder.states[max(0,len(recorder.states)-61)][cfg.N+1:])))
        else:
            atomic_npz(directory/'event_state.npz',time_s=np.array(event['time_s']),
                       state=recorder.states[recorder.kinds.index('event')],xi=grid.xi)
            nt=np.array([x[0] for x in neighborhood]);ny=np.array([x[1] for x in neighborhood])
            atomic_npz(directory/'event_neighborhood.npz',time_s=nt,offset_s=nt-event['time_s'],
                T_C=ny[:,:cfg.N+1],C=ny[:,cfg.N+1:],R_m=radius.at(nt),
                valid_interval_s=np.array([[x[2],x[3]] for x in neighborhood]),kind=np.array([x[4] for x in neighborhood]))
        recorder.flush(t,y,manifest,event)
    except BaseException as error:
        manifest.update(status='failed',failure_reason=repr(error),elapsed_this_call_s=perf_counter()-started)
        recorder.flush(t,y,manifest,event)
        raise
    return manifest


def offline_balance(directory):
    """只读原始场与表面探针，独立重算移动平衡，不调用RHS/在线残差。"""
    directory = Path(directory)
    cfg = Config.read(directory/'config.json')
    with np.load(directory/'fields.npz') as a:
        ts,xi,R,C = (a[x].copy() for x in ('time_s','xi','R_m','C'))
    with np.load(directory/'balance_probes.npz') as a:
        p = a['probes']
    face = np.r_[0.,(xi[:-1]+xi[1:])/2,1.]
    initial_vol = np.pi*cfg.L_m*cfg.R0_m**2*np.diff(face**2)
    inventory_scaled = C@initial_vol
    initial = float(inventory_scaled[0])
    if len(p):
        times,Cs,Rp,Ce = (p[:,:,i] for i in range(4))
        flux = 2*np.pi*cfg.L_m*Rp*cfg.hm*(Cs-Ce)/(Rp/cfg.R0_m)**2
        fine = simpson(flux,x=times,axis=1)
        coarse = simpson(flux[:,::2],x=times[:,::2],axis=1)
        ends = times[:,-1]
        if np.max(np.abs(times[1:,0]-ends[:-1]),initial=0) > 1e-7:
            raise ValueError('平衡探针区间有重叠或缺口')
        idx = np.searchsorted(ends,ts[1:])
        if np.any(idx >= len(ends)) or np.max(np.abs(ends[idx]-ts[1:]),initial=0) > 1e-7:
            raise ValueError('原始场采样没有对应的平衡积分端点')
        integral_fine = np.r_[0.,np.cumsum(fine)[idx]]
        integral_coarse = np.r_[0.,np.cumsum(coarse)[idx]]
    else:
        integral_fine=integral_coarse=np.zeros_like(ts)
    residual = (inventory_scaled-initial+integral_fine)/initial
    quadrature = (integral_fine-integral_coarse)/initial
    result = dict(status='PASS' if np.max(np.abs(residual)) <= 1e-6 and np.max(np.abs(quadrature)) <= 1e-7 else 'FAIL',
        max_balance_residual=float(np.max(np.abs(residual))),quadrature_difference=float(np.max(np.abs(quadrature))),
        probe_count=cfg.balance_probe_count,raw_rows=len(ts),probe_intervals=len(p),
        formula='(I_C/J-I_C(0)+integral(F_surface/J dt))/I_C(0)',
        fields_sha256=file_hash(directory/'fields.npz'),probes_sha256=file_hash(directory/'balance_probes.npz'))
    atomic_npz(directory/'balance_audit_series.npz',time_s=ts,inventory_over_J=inventory_scaled,
        integral_fine=integral_fine,integral_coarse=integral_coarse,residual=residual,quadrature=quadrature)
    save_json(directory/'balance_audit.json',result)
    return result

# === Q4 CORE END ===


def iterate_chunks(directory):
    directory=Path(directory)
    for chunk in read_json(directory/'raw/index.json')['chunks']:
        path=directory/chunk['path']
        if file_hash(path)!=chunk['sha256']:raise ValueError('原始分块哈希不一致')
        with np.load(path) as data:yield {name:data[name].copy() for name in data.files}


def load_fields(directory):
    chunks=list(iterate_chunks(directory))
    result={'xi':chunks[0]['xi']}
    for name in ('time_s','R_m','T_C','C','kind'):result[name]=np.concatenate([c[name] for c in chunks],axis=0)
    return result


def export_results(directory,output_directory,template=None):
    """从保存的真实计算值导出；固定物理位置在域外为空，表面单独取末节点。"""
    from openpyxl.styles import Font,Alignment,PatternFill
    directory,output_directory = Path(directory),Path(output_directory)
    output_directory.mkdir(parents=True,exist_ok=True)
    data = load_fields(directory)
    event = read_json(directory/'event.json')
    manifest = read_json(directory/'manifest.json')
    cfg = Config.read(directory/'config.json')
    if manifest['status']!='completed' or manifest['core_sha256'] != core_hash():
        raise ValueError('仅可导出当前核心完成的运行')
    root_t = event['time_s']; t=data['time_s']
    regular = np.arange(60.,root_t,60.)
    regular = regular[np.abs(regular-root_t)>1e-7]
    targets = np.r_[regular,root_t]
    def indices(times):
        result=[]
        for ti in times:
            j=int(np.argmin(np.abs(t-ti)))
            if abs(t[j]-ti)>1e-7:
                raise ValueError(f'原始场缺少规定时刻：{ti}s')
            result.append(j)
        return np.asarray(result,dtype=int)
    def values_for(times,spacing):
        ix=indices(times); radii=data['R_m'][ix]
        positions=np.arange(int(np.floor((max(radii)+1e-12)/spacing))+1)*spacing
        values=[]; masks=[]
        for j,R in zip(ix,radii):
            vals,mask=reconstruct_physical(data['C'][j],data['xi'],R,positions)
            values.append(np.r_[vals,data['C'][j,-1]])
            masks.append(mask)
        return ix,positions,np.array(values),np.asarray(masks,dtype=bool)
    ix,positions,values,mask = values_for(targets,.001)
    template = Path(template) if template else directory.parents[1]/'inputs/result4_template.xlsx'
    wb = load_workbook(template) if template.exists() else Workbook()
    ws = wb.active
    if ws.max_row: ws.delete_rows(1,ws.max_row)
    if ws.max_column: ws.delete_cols(1,ws.max_column)
    headers=['时间\\到药材中心的距离',*[float(x*100) for x in positions],'药材表面']
    ws.append(headers)
    for ti,vals in zip(targets,values):
        ws.append([float(ti),*[None if not np.isfinite(v) else float(v) for v in vals]])
    ws.freeze_panes='B2'
    ws.column_dimensions['A'].width=27
    for cell in ws[1]:
        cell.font=Font(name='宋体',bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='214A59')
        cell.alignment=Alignment(horizontal='center')
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.number_format='0.0000'
    for name in ('说明','表面半径'):
        if name in wb.sheetnames: del wb[name]
    notes=wb.create_sheet('说明')
    entries=[('项目','说明'),('run_id',manifest['run_id']),('模型','附录4；材料坐标；实测半径驱动；长度固定'),
        ('时间单位','s；每60s且严格早于临界根，末行另存未舍入临界根'),
        ('距离单位','cm；固定物理位置，不能将材料编号当作距离'),
        ('含水率单位','kg/kg（有效干基含水率）'),('域外单元格','空白表示该位置已经在药材之外，不等于0'),
        ('药材表面','直接取该时刻材料节点xi=1；真实半径见表面半径工作表'),
        ('固定位置恰等于半径','固定距离列与药材表面列可以取相同值，这不是重复计算错误'),
        ('临界根与严格小于','根上最大含水率等于0.15；严格达标采用根后真实积分和数值余量检查'),
        ('core_sha256',manifest['core_sha256']),('config_sha256',manifest['config_sha256']),
        ('input_sha256',manifest['input_sha256']),('tcrit_s',root_t),('tcrit_h',root_t/3600)]
    for row in entries: notes.append(row)
    notes.column_dimensions['A'].width=25; notes.column_dimensions['B'].width=100
    surface=wb.create_sheet('表面半径'); surface.append(['time_s','time_h','R_cm','C_surface'])
    for j,ti in zip(ix,targets):
        surface.append([float(ti),float(ti/3600),float(data['R_m'][j]*100),float(data['C'][j,-1])])
    path=output_directory/'result4.xlsx'; wb.save(path); wb.close()
    atomic_npz(output_directory/'result4_unrounded.npz',time_s=targets,time_h=targets/3600,
        physical_r_m=positions,R_m=data['R_m'][ix],C=values,inside=mask)
    atomic_npz(output_directory/'initial_state.npz',time_s=t[0],xi=data['xi'],R_m=data['R_m'][0],
        T_C=data['T_C'][0],C=data['C'][0])
    table_times=np.arange(21600.,root_t,21600.)
    table_times=np.r_[table_times[np.abs(table_times-root_t)>1e-7],root_t]
    tidx,tpos,tval,tmask=values_for(table_times,.005)
    tables=output_directory/'tables'; tables.mkdir(exist_ok=True)
    atomic_npz(tables/'table6_unrounded.npz',time_s=table_times,time_h=table_times/3600,
        physical_r_m=tpos,R_m=data['R_m'][tidx],C=tval,inside=tmask)
    with (tables/'table6.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f); writer.writerow(['time_s','time_h',*[f'r={x*100:g}cm' for x in tpos],'surface'])
        for ti,vals in zip(table_times,tval): writer.writerow([ti,ti/3600,*['' if not np.isfinite(v) else v for v in vals]])
    lines=['| 时刻/h | '+ ' | '.join([*[f'r={x*100:g} cm' for x in tpos],'药材表面'])+' |',
           '| --- | '+ ' | '.join(['---']*(len(tpos)+1))+' |']
    for i,(ti,vals) in enumerate(zip(table_times,tval)):
        label=f'{ti/3600:.4f}' if i<len(table_times)-1 else f'烘干结束时间（{ti/3600:.4f}）'
        lines.append('| '+label+' | '+' | '.join('—' if not np.isfinite(v) else f'{v:.4f}' for v in vals)+' |')
    (tables/'table6.md').write_text('\n'.join(lines)+'\n\n注：—表示域外；临界根为等号时刻，严格小于的后验见事件报告。\n',encoding='utf-8')
    # 逐格重读，不只看文件能否打开；包括空白掩膜、表面、根去重和有效列并集。
    check=load_workbook(path,read_only=True,data_only=True); saved=list(check[ws.title].values)
    assert len(saved)==len(targets)+1 and len(saved[0])==len(headers)
    max_error=0.; outside_filled=0; inside_missing=0
    assert np.all(np.diff(targets)>0) and targets[-1]==root_t
    for i,row in enumerate(saved[1:]):
        assert abs(row[0]-targets[i])<=1e-9
        for j,expected in enumerate(values[i]):
            actual=row[j+1]
            if np.isfinite(expected):
                inside_missing+=int(actual is None)
                assert actual is not None
                max_error=max(max_error,abs(actual-expected))
            else:
                outside_filled+=int(actual is not None)
    sr=list(check['表面半径'].values)[1:]
    for i,row in enumerate(sr):
        assert abs(row[0]-targets[i])<=1e-9 and abs(row[2]/100-data['R_m'][ix[i]])<=1e-12
    check.close()
    assert max_error<=1e-12 and outside_filled==0 and inside_missing==0
    audit=dict(status='PASS',run_id=manifest['run_id'],rows=len(targets),fixed_columns_cm=(positions*100).tolist(),
        table6_rows=len(table_times),xlsx_numeric_error=max_error,inside_missing=inside_missing,outside_filled=outside_filled,
        domain_outside_count=int((~mask).sum()),source_fields_sha256=file_hash(directory/'fields.npz'),
        result4_sha256=file_hash(path),template_sha256=file_hash(template) if template.exists() else None)
    save_json(output_directory/'export_check.json',audit)
    return audit


def main(argv=None):
    if hasattr(sys.stdout,"reconfigure"): sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description='第四问单文件核心与输入入口')
    parser.add_argument('action',nargs='?',default='run',choices=['inspect','prepare','run','balance','export'])
    parser.add_argument('--config',type=Path)
    parser.add_argument('--run-id')
    parser.add_argument('--resume',action='store_true')
    parser.add_argument('--output',type=Path,default=ROOT/'outputs/q4')
    args = parser.parse_args(argv)
    if args.action == 'prepare':
        result = prepare_inputs()
        update_stage('P00','running',['inputs/input_manifest.json'], '运行基础验证')
        print(json.dumps(result,ensure_ascii=False,indent=2))
    elif args.action == 'run':
        if args.resume and not args.run_id:
            parser.error('恢复需要指定已有--run-id')
        cfg = Config.read(args.config) if args.config else Config(**TEAM_PARAMETERS).validate()
        run_id = args.run_id or datetime.now().strftime('q4_%Y%m%d_%H%M%S')
        result = run_model(cfg,run_id,resume=args.resume)
        extensions = 0
        while result['end_reason']=='not_reached' and cfg.stop_at_event and extensions<20:
            if not result['extension_assessment']['environment_below_threshold']:
                break
            cfg = replace(cfg,end_s=cfg.end_s+cfg.extension_s)
            result = run_model(cfg,run_id,resume=True)
            extensions += 1
        print(json.dumps(result,ensure_ascii=False,indent=2))
        print('运行编号:',run_id)
        event_file = ROOT/'runs'/run_id/'event.json'
        if event_file.exists():
            ev=read_json(event_file)
            print(f'临界时间: {ev["time_h"]:.10f} h = {ev["time_s"]:.6f} s')
            print('导出命令: python 问题四_核心求解_单文件版.py export --run-id '+run_id)
    elif args.action == 'balance':
        if not args.run_id:
            parser.error('balance需要--run-id')
        print(json.dumps(offline_balance(ROOT/'runs'/args.run_id),ensure_ascii=False,indent=2))
    elif args.action == 'export':
        if not args.run_id:
            parser.error('export需要--run-id')
        print(json.dumps(export_results(ROOT/'runs'/args.run_id,args.output),ensure_ascii=False,indent=2))
    else:
        print(json.dumps(dict(root=str(ROOT),python=sys.executable,python_version=platform.python_version(),
            numpy=np.__version__,scipy=scipy.__version__,core_hash=core_hash(),
            existing_runs=[str(p) for p in (ROOT/'runs').glob('*/manifest.json')]),ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
