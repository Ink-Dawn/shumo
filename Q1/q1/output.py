"""结果导出"""
from dataclasses import asdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import csv
import json
import time

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .metrics import KEY_RADII_CM, KEY_TIMES, loss_depth
from .model import T_END
from .solver import sample


def round4(x):
    """十进制四舍五入到 4 位小数"""
    return float(Decimal(str(float(x))).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP))


def portable(path):
    """
    把数据来源写成相对包根的路径
    """
    p = Path(path)
    try:
        return p.resolve().relative_to(Path(__file__).resolve().parents[1]).as_posix()
    except ValueError:
        return p.name


def write_workbook(path, times, radii, fields, template=None):
    """
    生成结果表
    """
    path = Path(path)
    if template and path.resolve() == Path(template).resolve():
        raise ValueError('输出路径不能覆盖模板')
    wb = load_workbook(template) if template else Workbook()
    if not template:
        wb.remove(wb.active)

    for name, data in fields.items():
        ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        for merged in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(merged))
        for row in ws:
            for cell in row:
                cell.value = None

        ws.cell(1, 1, '时间/s \\ 到药材中心的距离/cm')
        for j, r in enumerate(radii, 2):
            ws.cell(1, j, round(r * 100, 8)).number_format = '0.0'
        for i, t in enumerate(times, 2):
            ws.cell(i, 1, int(t))
            for j, v in enumerate(data[i - 2], 2):
                ws.cell(i, j, round4(v)).number_format = '0.0000'

        ws.freeze_panes = 'B2'
        ws.column_dimensions['A'].width = 27
        ws.row_dimensions[1].height = 32
        for j in range(2, len(radii) + 2):
            ws.column_dimensions[get_column_letter(j)].width = 11
        for cell in ws[1]:
            cell.font = Font(name='微软雅黑', bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='245979')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.print_title_rows = '1:1'

    wb.save(path)
    wb.close()


def export_run(out, env, p, grid, heat, wet, template=None, flux='kirchhoff'):
    """导出 result1.xlsx、全精度 npz、题定节点 CSV 与运行记录。"""
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)

    times = np.arange(1.0, T_END + 1.0)          # 1 s 到 1800 s 的整秒
    radii = np.linspace(0.0, p.R, 21)            # 0 到 2 cm，每 0.1 cm 一个
    T = sample(heat, grid, times, radii)
    C = sample(wet, grid, times, radii)

    write_workbook(out / 'result1.xlsx', times, radii, {'温度': T, '水分浓度': C}, template)
    np.savez_compressed(out / '全精度结果.npz', time=times, radius_m=radii, T=T, C=C)

    # 21 列里取 0、0.5、1、1.5、2 cm 五列，步长正好是 5
    for label, field in [('表1_温度', T), ('表2_水分浓度', C)]:
        with (out / f'{label}.csv').open('w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(['时间/s', *[f'{x:g}' for x in KEY_RADII_CM]])
            for t in KEY_TIMES:
                writer.writerow([int(t), *[f'{v:.4f}' for v in field[t - 1, ::5]]])

    summary = []
    for t in KEY_TIMES:
        cw = wet.at(t)
        depth = loss_depth(grid, cw, p.C0) * 1000.0
        summary.append({
            'time_s': int(t),
            'T_center': float(heat.at(t)[0]), 'T_surface': float(heat.at(t)[-1]),
            'C_center': float(cw[0]), 'C_surface': float(cw[-1]),
            'C_average': float(grid.V @ cw / grid.V.sum()),
            'depth_5pct_mm': None if not np.isfinite(depth) else float(depth),
        })

    meta = {
        'source': portable(env.source),
        'generated_at': time.strftime('%Y-%m-%d %H:%M:%S'),
        'parameters': asdict(p),
        'grid': {'N': len(grid.r) - 1, 'beta': grid.beta},
        'solver': {'flux': flux, 'temperature': heat.stats, 'moisture': wet.stats,
                   'end_s': T_END},
        'summary': summary,
    }
    (out / '运行记录.json').write_text(
        json.dumps(meta, ensure_ascii=False, indent=2), encoding='utf-8')
    return meta
